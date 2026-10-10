"""Xuất file import theo loại cột: cả hai (mặc định, y hệt trước) / chỉ TSKT / chỉ FILTER.
Chỉ lọc CỘT khi ghi file — giá trị từng ô không đổi. Chạy: python tests/test_xuat_loai_cot.py"""
import io, re, sys, zipfile
from pathlib import Path

sys.path[:0] = [str(Path(__file__).resolve().parent.parent)]
import pandas as pd
from openpyxl import load_workbook

import pim_core as C


def _bang():
    def r(sku, mo, **v):
        return {"sku": sku, "vals": {"model_code": mo, "variant_code": "", **v}}
    return {
        "100": {"title": "Loa", "attr": ["product_mass", "color_filter", "note_text"], "ten": {},
                "rows": [r("A", "M1", product_mass="5 kg", color_filter="1, 2", note_text="x"),   # có cả hai nhóm
                         r("B", "M2", product_mass="7"),                                          # chỉ TSKT
                         r("D", "M4", color_filter="3"),                                          # chỉ FILTER
                         r("C", "M3")]},                                                           # trống hoàn toàn
        "200": {"title": "Tủ", "attr": ["fridge_mass"], "ten": {},                               # không có cột FILTER
                "rows": [r("E", "M5", fridge_mass="9")]},
    }


def _doc(x):
    return {n: [[c.value for c in rr] for rr in load_workbook(io.BytesIO(d)).active.iter_rows()] for n, d, _ in x["files"]}


def _chay(lc, bo_trong=False, sua=None):
    imp = pd.DataFrame([["M1", "A", "", "100"]], columns=C.COT_IMPORT).iloc[0:0]
    return C.xuat_file_import(_bang(), imp, sua or {}, {}, {}, bo_dong_trong=bo_trong, loai_cot=lc)


def _tim(d, k):
    return next(v for n, v in d.items() if k in n)


def test_mac_dinh_la_ca_hai():
    x = _chay("ca_hai")
    d = _doc(x)
    assert all("_CHI_" not in n for n in d)  # tên file không đổi
    assert _tim(d, "Loa")[0] == ["model_code", "variant_code", "product_mass", "color_filter", "note_text"]
    assert x["khong_cot"] == []
    # không truyền loai_cot = ca_hai: nội dung từng file y hệt
    imp = pd.DataFrame(columns=C.COT_IMPORT)
    mac_dinh = C.xuat_file_import(_bang(), imp, {}, {}, {})
    assert list(_doc(mac_dinh).values()) == list(d.values()) and mac_dinh["loai_cot"] == "ca_hai"


def test_chi_tskt_va_chi_filter():
    t = _doc(_chay("tskt"))
    f = _doc(_chay("filter"))
    assert all("_CHI_TSKT_" in n for n in t) and all("_CHI_FILTER_" in n for n in f)
    assert _tim(t, "Loa")[0] == ["model_code", "variant_code", "product_mass", "note_text"]
    assert _tim(f, "Loa")[0] == ["model_code", "variant_code", "color_filter"]
    # ô FILTER: vẫn gọn "1,2" như trước
    assert [row for row in _tim(f, "Loa")[2:] if row[0] == "M1"][0][2] == "1,2"
    # giá trị TSKT y nguyên
    assert [row for row in _tim(t, "Loa")[2:] if row[0] == "M1"][0][2:] == ["5 kg", "x"]


def test_nganh_khong_co_cot_nhom_chon_bi_bo_qua():
    x = _chay("filter")
    assert x["khong_cot"] == ["Tủ"] and not any("Tủ" in n for n in _doc(x))
    assert _chay("tskt")["khong_cot"] == [] and any("Tủ" in n for n in _doc(_chay("tskt")))


def test_sku_trong_va_sku_chi_co_nhom_kia():
    # bo_dong_trong: SKU trống hẳn (C) bị tách ở MỌI chế độ; SKU chỉ có TSKT (B) KHÔNG bị coi là "không có data" khi xuất chỉ FILTER
    for lc in ("ca_hai", "tskt", "filter"):
        x = _chay(lc, bo_trong=True)
        mo = [row[0] for row in _tim(_doc(x), "Loa")[2:]]
        assert "M3" not in mo, (lc, mo)
        assert x.get("bo_trong") == 1, (lc, x.get("bo_trong"))
    f = [row[0] for row in _tim(_doc(_chay("filter", True)), "Loa")[2:]]
    assert sorted(f) == ["M1", "M2", "M4"]  # M2 chỉ có TSKT: giữ dòng (ô FILTER trống) như xuất đủ
    t = [row[0] for row in _tim(_doc(_chay("tskt", True)), "Loa")[2:]]
    assert sorted(t) == ["M1", "M2", "M4"]  # M4 chỉ có FILTER: giữ dòng
    c = [row[0] for row in _tim(_doc(_chay("ca_hai", True)), "Loa")[2:]]
    assert sorted(c) == ["M1", "M2", "M4"]


def test_sua_tay_chi_liet_ke_dung_nhom():
    sua = {("100", "A", "product_mass"): "6", ("100", "A", "color_filter"): "1 , 2"}
    def txt(lc):
        z = zipfile.ZipFile(io.BytesIO(_chay(lc, sua=sua)["zip"]))
        return z.read("_CHINH_SUA_TAY.txt").decode("utf-8-sig") if "_CHINH_SUA_TAY.txt" in z.namelist() else ""
    assert "product_mass" in txt("ca_hai") and "color_filter" in txt("ca_hai")
    assert "product_mass" in txt("tskt") and "color_filter" not in txt("tskt")
    assert "color_filter" in txt("filter") and "product_mass" not in txt("filter")


def test_loai_cot_sai():
    try:
        _chay("abc")
    except ValueError:
        return
    raise AssertionError("phải báo lỗi")


if __name__ == "__main__":
    for n, f in list(globals().items()):
        if n.startswith("test_"):
            f()
    print("OK — xuất theo loại cột")
