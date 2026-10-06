"""So desktop (66.py) với web trên FILE THẬT: python tests/so_sanh_desktop_file_that.py "file mẫu.xlsx"
(tạo ws_desk.xlsx + out_desk/ ở thư mục đang đứng)."""
import sys, time, io, contextlib, shutil
import os
_T = os.path.dirname(os.path.abspath(__file__))
sys.path[:0] = [os.path.dirname(_T), _T, os.path.join(_T, "desktop_goc")]
from test_parity_desktop import _stub_tk; _stub_tk()
import pim_app_66 as D, pim_core as W
from pathlib import Path
from openpyxl import load_workbook
_goc = D.tim_cot_theo_ten
ALIAS = {"Mã thuộc tính": ("MÃ THUỘC TÍNH TSKT",), "Tên TSKT (MASTER)": ("TÊN MASTER",)}
def tim(header, *ten):
    extra = tuple(a for t in ten for a in ALIAS.get(t, ()))
    return _goc(header, *(ten + extra))
D.tim_cot_theo_ten = tim
f = sys.argv[1]
shutil.copy(f, "ws_desk.xlsx"); out = Path("out_desk"); shutil.rmtree(out, ignore_errors=True)
t = time.time(); buf = io.StringIO()
with contextlib.redirect_stdout(buf):
    D.chay_tat_ca(Path("ws_desk.xlsx"), out, xuat=False)
x = D.xuat_file_import_nhanh(Path("ws_desk.xlsx"), out)
print("DESKTOP", round(time.time() - t, 1), "s", x.get("files"))
print("\n".join(l for l in buf.getvalue().splitlines() if l.startswith("B")))
def dump(b):
    ws = load_workbook(io.BytesIO(b)).active
    return [[("" if c.value is None else str(c.value)) for c in r] for r in ws.iter_rows()]
d = {p.name.rsplit("_", 2)[0]: dump(p.read_bytes()) for p in Path(x["thu_muc"]).glob("*.xlsx")}
t = time.time()
w = W.doc_workspace_cu(open(f, "rb").read(), Path(f).name)
r = W.chay_map(w["data_sp"], w["import"], w["cau_hinh"], w["map_tskt"], w["map_filter"], W.option_maps(w["data_pim"]))
xw = W.xuat_file_import(r["bang"], w["import"], {}, {}, {})
print("WEB", round(time.time() - t, 1), "s", [(a, c) for a, _, c in xw["files"]], r["tom_tat"])
wd = {n.rsplit("_", 2)[0]: dump(b) for n, b, _ in xw["files"]}
assert set(d) == set(wd), (set(d), set(wd))
loi = 0
for k in d:
    A, B = d[k], wd[k]
    if A[:2] != B[:2]:
        print("HEADER KHÁC", k); print(A[:2]); print(B[:2]); loi += 1
    for i, (ra, rb) in enumerate(zip(A[2:], B[2:])):
        if ra != rb:
            loi += 1
            if loi < 6:
                print("DÒNG", i, [(A[0][j], ra[j], rb[j]) for j in range(len(ra)) if j < len(rb) and ra[j] != rb[j]][:6])
    if len(A) != len(B): print("SỐ DÒNG KHÁC", k, len(A), len(B)); loi += 1
n_o = sum(1 for k in d for rr in d[k][2:] for v in rr[2:] if v)
print("Ô có giá trị:", n_o, "| KHÁC NHAU:", loi)
