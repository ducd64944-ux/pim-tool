# -*- coding: utf-8 -*-
"""
pim_core.py — ENGINE (không phụ thuộc Streamlit) của tool TGDĐ CMS -> PIM bản web.

Port lại logic đã chạy thật ở bản desktop 66.py, nhưng làm trên DataFrame /
dict trong bộ nhớ thay vì 1 file Excel workspace:

  DATA SP (CMS export) + IMPORT (model/sku/variant/category)
      + CẤU HÌNH CATEGORY + MAPPING TSKT + MAPPING FILTER + DATA PIM
          ──► chay_map()  ──► "bảng" từng ngành hàng (giá trị mới)
      + SPEC PIM cũ (file export PIM)
          ──► tinh_kiem_tra() ──► cảnh báo / khác spec / đơn vị / Không-Đang cập nhật
          ──► xuat_file_import() ──► .xlsx MODEL + BIẾN THỂ (ô ép Text)

Nguyên tắc nghiệp vụ giữ nguyên bản desktop:
  1) TSKT -> luôn là TEXT nguyên văn PROPVALUE, nhiều giá trị nối "|".
  2) FILTER -> luôn là MÃ option tra qua DATA PIM, nhiều mã nối ", ";
     không tra được -> để trống + ghi log (không bịa).
  3) Đơn vị chỉ thêm vào ô SỐ TRƠN; không đụng FILTER.
  4) Giá trị PIM dạng mảng ["905", " 903"] -> 905, 903.
"""
from __future__ import annotations

import io
import json
import re
import unicodedata
import zipfile
from collections import defaultdict, Counter
from datetime import datetime, timezone, timedelta
from functools import lru_cache
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

REF_DIR = Path(__file__).resolve().parent / "data" / "reference"
VN_TZ = timezone(timedelta(hours=7))


def bay_gio() -> str:
    return datetime.now(VN_TZ).strftime("%Y-%m-%d %H:%M:%S")


# ============================================================================
# §1 CHUẨN HOÁ
# ============================================================================
_WS_RE = re.compile(r"\s+")
_NBSP = " "
COT_CO_DINH = ["model_code", "sku", "variant_code"]
COT_CO_DINH_LOWER = {"model_code", "sku", "variant_code"}          # cột đầu file, không lặp trong cấu hình
COT_KHONG_DIEN = {"model_code", "sku", "category_code", "variant_code"}  # = CFG.cot_co_dinh bản desktop
COT_KHONG_PHAI_SPEC = {"model_code", "sku", "category_code", "variant_code", "family_code",
                       "family_variant_code", "model_activated", "variant_activated"}
SEP_TSKT = "|"
SEP_FILTER = ", "


def ep_text(v) -> str:
    """Mọi giá trị -> chuỗi; số nguyên dạng float (215361.0) -> '215361'."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        if v != v:  # NaN
            return ""
        return str(int(v)) if v == int(v) else str(v)
    s = str(v)
    return "" if s in ("nan", "NaN", "None", "<NA>") else s


# Các hàm chuẩn hoá chuỗi là hàm THUẦN (cùng đầu vào -> cùng kết quả) nên được nhớ đệm (lru_cache):
# dữ liệu lớn lặp lại rất nhiều giá trị (Có/Không, mã option, mã cột…) -> nhanh hơn nhiều lần, kết quả y nguyên.
_CACHE = 1 << 18
_CHU_RONG = frozenset(("nan", "NaN", "None", "<NA>"))

# KÝ TỰ ẨN hay gặp khi copy từ web/Word/CMS: không nhìn thấy nhưng làm SKU/giá trị so khớp SAI và lọt vào file import
# (PIM nhận cả ký tự ẩn -> ô trông giống nhau mà không trùng). Bỏ hẳn; khoảng trắng lạ (NBSP, en-space…) -> ' '.
# Giữ nguyên \n và \t (giá trị nhiều dòng). \r bị bỏ vì xlsx ghi \r thành "_x000D_" trên PIM.
_KY_TU_AN = "\u200b\u200c\u200d\u200e\u200f\u2060\u2061\u2062\u2063\u2064\ufeff\u00ad\u180e\u202a\u202b\u202c\u202d\u202e\r"
_KHOANG_TRANG_LA = "\u00a0\u2000\u2001\u2002\u2003\u2004\u2005\u2006\u2007\u2008\u2009\u200a\u202f\u205f\u3000\u2028\u2029"
_AN_MAP = {ord(c): None for c in _KY_TU_AN}
_AN_MAP.update({c: None for c in list(range(0x00, 0x09)) + [0x0b, 0x0c] + list(range(0x0e, 0x20)) + [0x7f]})
_AN_MAP.update({ord(c): " " for c in _KHOANG_TRANG_LA})
_AN_RE = re.compile("[" + re.escape(_KY_TU_AN) + "\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def bo_ky_tu_an(s: str) -> str:
    """Bỏ ký tự ẩn / ký tự điều khiển, đổi khoảng trắng lạ thành ' ' (không trim)."""
    return s.translate(_AN_MAP) if s else s


def co_ky_tu_an(s) -> bool:
    return bool(s) and isinstance(s, str) and (_AN_RE.search(s) is not None
                                               or any(c in s for c in _KHOANG_TRANG_LA))


@lru_cache(maxsize=_CACHE)
def _ck_str(s: str) -> str:
    return "" if s in _CHU_RONG else s.translate(_AN_MAP).strip()


def chuan_hoa_key(v) -> str:
    if type(v) is str:
        return _ck_str(v)
    return ep_text(v).translate(_AN_MAP).strip()


@lru_cache(maxsize=_CACHE)
def _code_str(s: str) -> str:
    return _WS_RE.sub("", _ck_str(s))


def chuan_hoa_code(v) -> str:
    if type(v) is str:
        return _code_str(v)
    return _WS_RE.sub("", chuan_hoa_key(v))


_ID_THAP_PHAN_RE = re.compile(r"\d+\.0+")


@lru_cache(maxsize=_CACHE)
def _id_str(s: str) -> str:
    if _ID_THAP_PHAN_RE.fullmatch(s):
        return s.split(".")[0]
    return s


def chuan_hoa_id(v) -> str:
    return _id_str(chuan_hoa_code(v))


@lru_cache(maxsize=_CACHE)
def chuan_hoa_ten(v) -> str:
    """So khớp TÊN thuộc tính: NFC, thường, gọn khoảng trắng, bỏ dấu câu cuối."""
    s = unicodedata.normalize("NFC", _WS_RE.sub(" ", chuan_hoa_key(v))).lower()
    return s.strip(" .:;-")


@lru_cache(maxsize=4096 * 4)
def la_cot_filter(code: str) -> bool:
    """Cột FILTER (nhận MÃ option). Desktop: có chữ "filter" trong mã. Sửa 1 nhầm lẫn: cột TSKT có chữ filter
    trong tên thuộc tính (vd filter_tskt_master = "Bộ lọc" của robot hút bụi) KHÔNG phải cột FILTER."""
    c = (code or "").lower()
    return "filter" in c and "tskt" not in c


# ============================================================================
# §2 ĐƠN VỊ / GIÁ TRỊ RỖNG / LÀM SẠCH GIÁ TRỊ PIM
# ============================================================================
DON_VI_LUA_CHON = ["", "cm", "mm", "m", "kg", "g", "inch"]
_SO_TRON_RE = re.compile(r"^\d+(?:[.,]\d+)?$")
_KT_CODE_RE = re.compile(
    r"(?:^|_)(deep|depth|high|height|horizontal|vertical|width|wide|length|long|thick|thickness|"
    r"mass|weight|size|dimension|dimensions|diameter|tall)(?:_|$)")
_KT_TEN_RE = re.compile(
    r"^(sâu|cao|ngang|dọc|dày|dài|rộng|nặng)\b|khối lượng|trọng lượng|kích thước|kích cỡ|đường kính",
    re.IGNORECASE)
_BO_DV_RE = re.compile(r"\s*(cm|mm|m|kg|g|inch)$", re.IGNORECASE)
_DV_DINH_RE = re.compile(r"(\d)\s+(cm|mm|m|kg|g|inch)\b")


@lru_cache(maxsize=_CACHE)
def _so_tron_str(s: str) -> bool:
    return bool(_SO_TRON_RE.fullmatch(s))


def la_so_tron(v) -> bool:
    return _so_tron_str(chuan_hoa_key(v))


def them_don_vi(v, dv: str):
    s, dv = chuan_hoa_key(v), chuan_hoa_key(dv)
    if not dv or not s or not _SO_TRON_RE.fullmatch(s):
        return v
    return f"{s} {dv}"


# ---- BIẾN ĐỔI HÀNG LOẠT theo cột (mở rộng của ĐƠN VỊ; không bao giờ áp vào cột FILTER) ----
BD_KIEU = {"them_sau": "Thêm vào SAU (đơn vị / chữ)", "them_truoc": "Thêm vào TRƯỚC",
           "ca_hai": "Thêm cả TRƯỚC và SAU",
           "doi_dv": "Đổi đơn vị (quy đổi số: mm→cm…)", "thay": "Thay chữ A → B", "lam_tron": "Làm tròn số (A = số lẻ)"}
BD_PHAM_VI = {"so": "Chỉ ô SỐ TRƠN (đúng rule cũ: 9 → 9 kg)", "tung_phan": "Từng giá trị số trong ô (9|10 → 9 kg|10 kg)",
              "tat_ca": "Mọi ô có giá trị"}
HE_SO_DV = {("mm", "cm"): 0.1, ("cm", "mm"): 10, ("m", "cm"): 100, ("cm", "m"): 0.01, ("mm", "m"): 0.001,
            ("m", "mm"): 1000, ("g", "kg"): 0.001, ("kg", "g"): 1000, ("inch", "cm"): 2.54, ("cm", "inch"): 1 / 2.54,
            ("ml", "lít"): 0.001, ("lít", "ml"): 1000, ("w", "kw"): 0.001, ("kw", "w"): 1000,
            ("mah", "ah"): 0.001, ("ah", "mah"): 1000, ("phút", "giờ"): 1 / 60, ("giờ", "phút"): 60}
_SO_PHAN_RE = re.compile(r"^(-?\d+(?:[.,]\d+)?)\s*([^\d\s].*)?$")


def _fmt_so(x: float) -> str:
    t = f"{x:.4f}".rstrip("0").rstrip(".")
    return t if t not in ("", "-0") else "0"


def _noi(v: str, them: str, truoc: bool) -> str:
    t = them.strip()
    if not t:
        return v
    if truoc:
        return v if v.startswith(t) else (t + ("" if t[-1] in "([/-" else " ") + v)
    return v if v.endswith(t) else (v + ("" if t[0] in ",.;:)/%" else " ") + t)


def ap_buoc(v: str, b: dict) -> str:
    """Áp 1 bước biến đổi cho 1 ô (chuỗi đã chuẩn hoá)."""
    if not v:
        return v
    kieu, pv = b.get("kieu"), b.get("pham_vi", "so")
    a, bb = str(b.get("a") or ""), str(b.get("b") or "")
    if kieu == "thay":
        return v.replace(a, bb) if a else v
    phan = v.split(SEP_TSKT) if pv == "tung_phan" else [v]

    def mot(p: str) -> str:
        p2 = p.strip()
        so = la_so_tron(p2)
        if kieu in ("them_sau", "them_truoc", "ca_hai"):
            if pv in ("so", "tung_phan") and not so:
                return p
            if kieu == "ca_hai":  # a = chữ phía trước, b = chữ phía sau
                return _noi(_noi(p2, bb, False), a, True)
            return _noi(p2, a, kieu == "them_truoc")
        m = _SO_PHAN_RE.match(p2)
        if not m:
            return p
        x = float(m.group(1).replace(",", "."))
        dv = (m.group(2) or "").strip()
        if kieu == "lam_tron":
            try:
                n = int(float(a or 0))
            except ValueError:
                n = 0
            return f"{_fmt_so(round(x, n))}{(' ' + dv) if dv else ''}"
        if kieu == "doi_dv":
            if dv and dv.lower() != a.strip().lower():
                return p
            if not dv and pv == "tat_ca" and not so:
                return p
            hs = HE_SO_DV.get((a.strip().lower(), bb.strip().lower()))
            if hs is None:
                try:
                    hs = float(str(b.get("he_so") or "").replace(",", "."))
                except ValueError:
                    return p
            return f"{_fmt_so(x * hs)} {bb.strip()}".strip()
        return p
    if pv == "so" and kieu in ("them_sau", "them_truoc", "ca_hai", "doi_dv", "lam_tron") and len(phan) == 1:
        return mot(phan[0])
    return SEP_TSKT.join(mot(p) for p in phan)


MAU_GOP_KT = {
    "{1} x {2} x {3} {dv}": "30 x 20 x 10 cm",
    "Dài {1} - Rộng {2} - Cao {3}": "Dài 30 cm - Rộng 20 cm - Cao 10 cm",
    "Ngang {1} - Cao {2} - Sâu {3}": "Ngang 30 cm - Cao 20 cm - Sâu 10 cm",
    "Dài {1} x Rộng {2} x Cao {3} {dv}": "Dài 30 x Rộng 20 x Cao 10 cm",
    "{1} x {2} {dv}": "30 x 20 cm (2 giá trị)",
}


def gop_kich_thuoc(gia_tri: List[str], mau: str, bo_dv: bool = True, dv: str = "", thieu_bo_qua: bool = True) -> str:
    """Gộp 2–4 giá trị (dài/rộng/cao…) thành 1 chuỗi theo mẫu có {1} {2} {3} {4} và {dv}.
    bo_dv=True: chỉ lấy SỐ của từng giá trị, đơn vị đặt 1 lần ở {dv} (lấy đơn vị đầu tiên gặp nếu dv trống).
    Thiếu 1 giá trị: thieu_bo_qua=True -> trả '' (không gộp)."""
    so, dv_gap = [], ""
    for v in gia_tri:
        v = chuan_hoa_key(v)
        if not v:
            if thieu_bo_qua:
                return ""
            so.append("")
            continue
        p = v.split(SEP_TSKT)[0].strip()
        m = _SO_PHAN_RE.match(p)
        if bo_dv and m:
            so.append(m.group(1))
            dv_gap = dv_gap or (m.group(2) or "").strip()
        else:
            so.append(p)
    don_vi = (dv or dv_gap).strip()
    out = mau
    for i in range(4):
        out = out.replace("{" + str(i + 1) + "}", so[i] if i < len(so) else "")
    if not bo_dv:
        out = out.replace("{dv}", "")
    elif "{dv}" in out:
        out = out.replace("{dv}", don_vi)
    elif don_vi:  # mẫu dạng "Dài {1} - Rộng {2}…": gắn đơn vị sau từng số
        for x in so:
            if x and re.fullmatch(r"-?\d+(?:[.,]\d+)?", x):
                out = re.sub(rf"(?<![\d.,]){re.escape(x)}(?![\d.,]|\s*{re.escape(don_vi)})", f"{x} {don_vi}", out, count=1)
    return _WS_RE.sub(" ", out).strip(" -x")


def ap_bien_doi(v: str, buoc: List[dict]) -> str:
    for b in buoc or []:
        v = ap_buoc(v, b)
    return v


def la_cot_kich_thuoc(code: str, ten: str) -> bool:
    if not code or la_cot_filter(code) or code.lower() in COT_KHONG_PHAI_SPEC:
        return False
    return bool(_KT_CODE_RE.search(code.lower()) or _KT_TEN_RE.search(chuan_hoa_key(ten)))


def goi_y_don_vi(code: str, ten: str) -> str:
    c, t = code.lower(), chuan_hoa_key(ten).lower()
    if re.search(r"khối lượng|trọng lượng|^nặng", t):
        return "kg"
    if re.search(r"(?:^|_)(mass|weight)(?:_|$)", c) and not re.search(r"(?:^|_)size(?:_|$)", c):
        return "kg"
    if "screen" in c or "màn hình" in t:
        return "inch"
    return "cm" if la_cot_kich_thuoc(code, ten) else ""


def chi_khac_don_vi(moi: str, old: str) -> bool:
    o, m = chuan_hoa_key(old), chuan_hoa_key(moi)
    if not o or not m or not la_so_tron(o) or not _BO_DV_RE.search(m):
        return False
    return _BO_DV_RE.sub("", m).replace(",", ".") == o.replace(",", ".")


_MANG_RE = re.compile(r"^\s*\[(.*)\]\s*$", re.S)


def lam_sach_gia_tri_pim(v, code: str = "") -> str:
    """["905", " 903"] -> '905, 903' (FILTER) / 'a|b' (TSKT)."""
    return _lam_sach(chuan_hoa_key(v), la_cot_filter(code))


@lru_cache(maxsize=_CACHE)
def _lam_sach(s: str, la_filter: bool) -> str:
    m = _MANG_RE.match(s)
    if not m:
        return s
    trong = m.group(1)
    items = re.findall(r'"((?:[^"\\]|\\.)*)"', trong) or trong.split(",")
    items = [chuan_hoa_key(x).strip("'") for x in items]
    items = [x for x in items if x]
    return (SEP_FILTER if la_filter else SEP_TSKT).join(items)


TRANG_THAI_GIONG = "Giống"
TRANG_THAI_KHAC = "KHÁC GIÁ TRỊ"
TRANG_THAI_TOOL_TRONG = "TOOL ĐỂ TRỐNG — PIM ĐANG CÓ"
TRANG_THAI_PIM_TRONG = "PIM TRỐNG — TOOL THÊM MỚI"
TRANG_THAI_DON_VI = "CHỈ THÊM ĐƠN VỊ"
TRANG_THAI_BO_QUA = "ĐỂ TRỐNG — KHÔNG CẬP NHẬT"
GIA_TRI_RONG_MAC_DINH = ["Không", "Không có", "Đang cập nhật", "Hãng không công bố", "Không công bố",
                         "Chưa có thông tin", "Chưa cập nhật", "Không rõ", "-", "N/A"]
HD_GIU, HD_TRONG, HD_THAY = "giu", "trong", "thay"
NHAN_HD = {HD_GIU: "Giữ — import như cũ", HD_TRONG: "Để trống — không cập nhật", HD_THAY: "Thay bằng…"}


@lru_cache(maxsize=_CACHE)
def _khoa_rong_str(s: str) -> str:
    return unicodedata.normalize("NFC", _WS_RE.sub(" ", s)).strip(" .;:").lower()


def khoa_gia_tri_rong(v) -> str:
    return _khoa_rong_str(chuan_hoa_key(v))


_KHOA_RONG_MAC_DINH = {khoa_gia_tri_rong(x) for x in GIA_TRI_RONG_MAC_DINH}


def ap_quy_tac_rong(v: str, code: str, rong: Optional[Dict[str, list]]) -> Tuple[str, Optional[str]]:
    if not rong or not v or la_cot_filter(code) or code.lower() in COT_KHONG_PHAI_SPEC:
        return v, None
    q = rong.get(khoa_gia_tri_rong(v))
    if not q or q[0] == HD_GIU:
        return v, None
    if q[0] == HD_TRONG:
        return "", HD_TRONG
    return chuan_hoa_key(q[1] if len(q) > 1 else ""), HD_THAY


@lru_cache(maxsize=_CACHE)
def khoa_option(v: str) -> str:
    """Khoá tra option FILTER: như desktop (trim + bỏ NBSP + chữ thường) và thêm NFC
    (chữ Việt dựng sẵn/tổ hợp coi là một) — chỉ làm KHỚP NHIỀU HƠN, không đổi mã đã khớp."""
    return unicodedata.normalize("NFC", v or "").lower()


def tap_gia_tri(v: str, code: str, option_map: Dict[Tuple[str, str], str]) -> frozenset:
    s = unicodedata.normalize("NFC", _WS_RE.sub(" ", lam_sach_gia_tri_pim(v, code)))
    if not s:
        return frozenset()
    if la_cot_filter(code):
        out = set()
        for t in re.split(r"[,|]", s):
            t = t.strip()
            if not t:
                continue
            tid = chuan_hoa_id(t)
            if not re.fullmatch(r"\d+", tid):
                tid = option_map.get((code, khoa_option(t)), tid)
            out.add(tid.lower())
        return frozenset(out)
    return frozenset(_DV_DINH_RE.sub(r"\1\2", p.strip().lower()) for p in s.split("|") if p.strip())


def ma_filter_tu_gia_tri(code: str, gia_tri: str, option_map: Dict[Tuple[str, str], str]) -> str:
    """Chuẩn hoá ô FILTER thành danh sách MÃ option (chữ -> mã qua DATA PIM)."""
    out: List[str] = []
    for t in re.split(r"[,|]", lam_sach_gia_tri_pim(gia_tri, code)):
        t = t.strip()
        if not t:
            continue
        tid = chuan_hoa_id(t)
        if not re.fullmatch(r"\d+", tid):
            tid = option_map.get((code, khoa_option(t)), t)
        if tid not in out:
            out.append(tid)
    return SEP_FILTER.join(out)


# ============================================================================
# §3 ĐỌC FILE (upload) -> DataFrame
# ============================================================================
COT_DATA_SP = ["PRODUCTID", "PRODUCTCODE", "PRODUCTNAME", "PROPERTYID", "PROPERTYNAME", "PROPVALUE", "CATEGORYID"]
COT_IMPORT = ["model_code", "sku", "variant_code", "category_code"]
COT_SPEC = ["sku", "model_code", "category_code", "variant_code", "ma", "ten", "gia_tri"]
COT_MAP_TSKT = ["cate", "cate_name", "prop_id", "prop_name", "ma", "ten_ma"]
COT_MAP_FILTER = ["cate", "cate_name", "prop_id", "prop_name", "ma"]
COT_DATA_PIM = ["Code", "Name", "Type", "Group", "Active", "OptionCode", "OptionValue"]


class SoExcel:
    """Đọc file Excel 1 LẦN, lấy nhiều sheet. Ưu tiên python-calamine (Rust — nhanh 20-50 lần openpyxl, không bị
    ảnh hưởng bởi vùng định dạng thừa hàng trăm nghìn dòng trống như file mẫu), tự lùi về openpyxl nếu thiếu/ lỗi.
    Giá trị ô trả về giống openpyxl data_only (số, chữ, ngày, True/False); ô trống = "" hoặc None."""

    def __init__(self, data: bytes, ten_file: str = "x.xlsx"):
        self.data, self.ten_file, self._cal, self.loai = data, ten_file, None, "openpyxl"
        if ten_file.lower().endswith((".csv", ".txt")):
            self.loai, self.sheets = "csv", ["CSV"]
            return
        try:
            from python_calamine import CalamineWorkbook
            self._cal = CalamineWorkbook.from_filelike(io.BytesIO(data))
            self.sheets, self.loai = list(self._cal.sheet_names), "calamine"
        except Exception:  # noqa: BLE001 - chưa cài / file lạ -> openpyxl
            wb = load_workbook(io.BytesIO(data), read_only=True)
            try:
                self.sheets = list(wb.sheetnames)
            finally:
                wb.close()

    def rows(self, sheet=0) -> List[list]:
        if self.loai == "csv":
            txt = self.data.decode("utf-8-sig", errors="replace")
            import csv
            try:
                dialect = csv.Sniffer().sniff(txt[:5000], delimiters=",;\t") if txt else csv.excel
            except csv.Error:
                dialect = csv.excel
            return [list(r) for r in csv.reader(io.StringIO(txt), dialect)]
        ten = self.sheets[sheet] if isinstance(sheet, int) else sheet
        if self._cal is not None:
            try:
                out = self._cal.get_sheet_by_name(ten).to_python(skip_empty_area=False)
                while out and not any(v not in ("", None) for v in out[-1]):
                    out.pop()
                return out
            except Exception:  # noqa: BLE001
                pass
        wb = load_workbook(io.BytesIO(self.data), read_only=True, data_only=True)
        try:
            sh = wb[ten]
            sh.reset_dimensions()  # bỏ qua <dimension> sai -> đọc đủ mọi ô
            out = [list(r) for r in sh.iter_rows(values_only=True)]
            while out and not any(v not in ("", None) for v in out[-1]):
                out.pop()
            return out
        finally:
            wb.close()


def _doc_bang_tho(data: bytes, ten_file: str, sheet=0) -> List[list]:
    """Đọc 1 sheet (xlsx/xlsm/xls/csv) ra list các dòng (giá trị thô)."""
    return SoExcel(data, ten_file).rows(sheet)


def _ten_sheet(data: bytes) -> List[str]:
    return SoExcel(data).sheets


@lru_cache(maxsize=4096)
def khoa_tieu_de(h) -> str:
    """So khớp TÊN CỘT: không phân biệt hoa/thường, dấu tiếng Việt, khoảng trắng, gạch dưới, ký tự lạ.
    "Mã sản phẩm ERP" = "ma_san_pham_erp" = "MÃ SẢN PHẨM ERP"."""
    s = unicodedata.normalize("NFD", chuan_hoa_key(h).lower().replace("đ", "d"))
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]", "", s)


# Tên cột chấp nhận cho từng trường (thêm tên mới vào đây là mọi chỗ đọc file đều nhận)
TEN_COT: Dict[str, List[str]] = {
    "sku": ["sku", "Mã sản phẩm ERP", "Mã ERP", "ERP", "Mã SP ERP", "Mã SP", "Mã sản phẩm", "Mã hàng", "productcode"],
    "model_code": ["model_code", "Mã model", "Model", "Mã mẫu", "modelcode"],
    "variant_code": ["variant_code", "Mã biến thể", "Biến thể", "variant", "variantcode"],
    "family_variant_code": ["family_variant_code", "Mã họ biến thể", "Họ biến thể", "familyvariantcode"],
    "category_code": ["category_code", "Mã danh mục PIM", "Danh mục PIM", "Mã danh mục", "categorycode"],
    "PRODUCTID": ["PRODUCTID", "product id", "Mã sản phẩm CMS", "ID sản phẩm"],
    "PRODUCTCODE": ["PRODUCTCODE", "product code", "Mã ERP", "Mã sản phẩm ERP", "sku"],
    "PRODUCTNAME": ["PRODUCTNAME", "product name", "Tên sản phẩm", "Tên SP"],
    "PROPERTYID": ["PROPERTYID", "property id", "Mã thuộc tính", "Mã thuộc tính CMS", "ID thuộc tính"],
    "PROPERTYNAME": ["PROPERTYNAME", "property name", "Tên thuộc tính", "Tên thuộc tính CMS"],
    "PROPVALUE": ["PROPVALUE", "prop value", "PROPERTYVALUE", "Giá trị", "Giá trị thuộc tính", "value"],
    "CATEGORYID": ["CATEGORYID", "category id", "Mã ngành hàng", "Mã ngành hàng CMS", "Mã ngành", "ID ngành hàng"],
}


def _tim_cot(header: List[str], *ten: str) -> int:
    """Vị trí (0-based) cột đầu tiên khớp 1 trong các tên (so theo khoa_tieu_de), -1 nếu không có.
    Ưu tiên theo thứ tự tên truyền vào."""
    kh = [khoa_tieu_de(h) for h in header]
    for t in ten:
        k = khoa_tieu_de(t)
        if k and k in kh:
            return kh.index(k)
    return -1


def tim_cot_truong(header: List[str], truong: str) -> int:
    return _tim_cot(header, *TEN_COT.get(truong, [truong]))


def doc_cms_export(data: bytes, ten_file: str) -> Tuple[pd.DataFrame, Optional[str]]:
    """File CMS export (cột PRODUCTID/PRODUCTCODE/.../PROPVALUE/CATEGORYID,
    nhận diện theo TÊN, không phụ thuộc thứ tự) -> DataFrame DATA SP."""
    so = SoExcel(data, ten_file)
    # file mẫu nhiều sheet (IMPORT đứng đầu): ưu tiên sheet "DATA SP", không thì sheet đầu tiên có cột PRODUCTCODE
    thu = (["DATA SP"] if "DATA SP" in so.sheets else []) + [x for x in so.sheets if x != "DATA SP"]
    rows = []
    for ten in thu:
        rows = so.rows(ten)
        if rows and _tim_cot([ep_text(h) for h in rows[0]], "PRODUCTCODE") >= 0:
            break
    if not rows:
        return pd.DataFrame(columns=COT_DATA_SP), "File trống."
    header = [ep_text(h) for h in rows[0]]
    idx = {c: tim_cot_truong(header, c) for c in COT_DATA_SP}
    thieu = [c for c, i in idx.items() if i < 0 and c not in ("PRODUCTID", "PRODUCTNAME", "CATEGORYID")]
    if thieu and all(idx[c] < 0 for c in ("PRODUCTCODE", "PROPERTYID", "PROPVALUE")) and len(rows[0]) >= 6:
        # Tiêu đề lạ / không có tiêu đề -> đọc THEO VỊ TRÍ như bản desktop (CFG.sp_*): A PRODUCTID, B PRODUCTCODE,
        # C PRODUCTNAME, D PROPERTYID, E PROPERTYNAME, F PROPVALUE, G CATEGORYID
        idx = {c: i for i, c in enumerate(COT_DATA_SP)}
        if re.fullmatch(r"\d+", chuan_hoa_id(rows[0][3])):  # dòng 1 đã là dữ liệu (không có tiêu đề)
            rows = [[]] + rows
        thieu = []
    if thieu:
        return pd.DataFrame(columns=COT_DATA_SP), (f"Thiếu cột: {', '.join(thieu)} (tên cột chấp nhận: "
                                                   + "; ".join(f"{c} = {' / '.join(TEN_COT[c][:4])}" for c in thieu) + ")")
    out = []
    for r in rows[1:]:
        if not r:
            continue
        def g(c):
            i = idx[c]
            return r[i] if 0 <= i < len(r) else None
        code = chuan_hoa_code(g("PRODUCTCODE"))
        if not code:
            continue
        out.append([chuan_hoa_id(g("PRODUCTID")), code, chuan_hoa_key(g("PRODUCTNAME")),
                    chuan_hoa_id(g("PROPERTYID")), chuan_hoa_key(g("PROPERTYNAME")),
                    chuan_hoa_key(g("PROPVALUE")), chuan_hoa_id(g("CATEGORYID"))])
    if not out:
        return pd.DataFrame(columns=COT_DATA_SP), "Không có dòng nào có PRODUCTCODE."
    return pd.DataFrame(out, columns=COT_DATA_SP), None


def la_bien_the_mau(ma_ho_bien_the) -> bool:
    """RULE: Mã họ biến thể có chữ "color" (vd lvl_1_color_iden_master) = biến thể THEO MÀU
    -> SKU đó xuất vào file import MODEL, không phải file BIENTHE (nên bỏ variant_code khi nạp)."""
    return "color" in chuan_hoa_key(ma_ho_bien_the).lower()


def doc_pim_export(data: bytes, ten_file: str) -> dict:
    """File export PIM (dòng 1 = mã cột, dòng 2 = tên, dữ liệu từ dòng 3):
    -> IMPORT (model/sku/variant/category) + SPEC (dạng dọc, đã làm sạch [""])."""
    rows = _doc_bang_tho(data, ten_file)
    if len(rows) < 2:
        return {"loi": "File không đủ 2 dòng tiêu đề (mã cột + tên cột)."}
    h1 = [chuan_hoa_key(h) for h in rows[0]]
    h2 = [chuan_hoa_key(h) for h in rows[1]] + [""] * len(h1)
    idx = {h.lower(): j for j, h in enumerate(h1) if h}
    for f in ("sku", "model_code", "variant_code", "category_code"):  # tên cột khác (Mã sản phẩm ERP…) cũng nhận
        if f not in idx:
            j = tim_cot_truong(h1, f)
            if j < 0:
                j = tim_cot_truong(h2, f)
            if j >= 0:
                idx[f] = j
    if "sku" not in idx:
        return {"loi": 'Không thấy cột sku (sku / Mã sản phẩm ERP / Mã ERP…) ở dòng 1-2 — cần file export PIM.'}
    if "family_variant_code" not in idx:
        j = tim_cot_truong(h1, "family_variant_code")
        if j < 0:
            j = tim_cot_truong(h2, "family_variant_code")
        if j >= 0:
            idx["family_variant_code"] = j
    co_dinh = {idx.get(f) for f in ("sku", "model_code", "variant_code", "category_code")}
    n_mau = 0
    spec_cols = [(h, j, h2[j]) for j, h in enumerate(h1) if h and h.lower() not in COT_KHONG_PHAI_SPEC
                 and j not in co_dinh and re.fullmatch(r"[A-Za-z0-9_]+", h) and "_" in h]
    imp, spec, seen = [], [], set()
    bo_khong_sku = trung = 0
    for r in rows[2:]:
        if not r or all(v in (None, "") for v in r):
            continue
        def g(k):
            j = idx.get(k)
            return chuan_hoa_key(r[j]) if j is not None and j < len(r) else ""
        sku = chuan_hoa_code(g("sku"))
        if not sku:
            bo_khong_sku += 1
            continue
        if sku in seen:
            trung += 1
            continue
        seen.add(sku)
        model, variant, cate = g("model_code"), g("variant_code"), chuan_hoa_id(g("category_code"))
        if variant and la_bien_the_mau(g("family_variant_code")):
            variant, n_mau = "", n_mau + 1
        imp.append([model, sku, variant, cate])
        co = False
        for code, j, ten in spec_cols:
            v = lam_sach_gia_tri_pim(r[j] if j < len(r) else "", code)
            if v:
                spec.append([sku, model, cate, variant, code, ten, v])
                co = True
        if not co:
            spec.append([sku, model, cate, variant, "", "", ""])
    if not imp:
        return {"loi": "Không có dòng dữ liệu nào có sku (từ dòng 3)."}
    df_imp = pd.DataFrame(imp, columns=COT_IMPORT)
    return {"loi": None, "import": df_imp, "spec": pd.DataFrame(spec, columns=COT_SPEC),
            "bo_khong_sku": bo_khong_sku, "trung": trung,
            "thieu_model": int((df_imp.model_code == "").sum()),
            "thieu_cate": int((df_imp.category_code == "").sum()), "so_cot_spec": len(spec_cols),
            "bien_the_mau": n_mau}


_MA_RE = re.compile(r"^[A-Za-z0-9._\-]+$")
_MA_COT_RE = re.compile(r"^[A-Za-z0-9]+(?:_[A-Za-z0-9]+)+$")  # mã thuộc tính PIM: abc_tskt_master, x_filter_master


def _la_dong_ten_vn(row: list, j: int) -> bool:
    """Dòng 2 kiểu desktop (tên tiếng Việt): ô ở cột sku trống hoặc có dấu/khoảng trắng."""
    v = chuan_hoa_key(row[j]) if 0 <= j < len(row) else ""
    return not v or bool(re.search(r"[^\x00-\x7F]|\s", v))


def doc_mot_cuc(rows: List[list]) -> dict:
    """NẠP 1 CỤC: bảng bất kỳ có model / SKU / biến thể (+ category, + các cột TSKT/FILTER) — có hay không có
    tiêu đề, tên cột tiếng Việt/Anh, thứ tự tuỳ ý (dán từ Excel cũng được).
      * Có tiêu đề: nhận cột theo TEN_COT; dòng 2 là tên tiếng Việt thì bỏ qua (như desktop: dữ liệu từ dòng 3).
      * Không tiêu đề: đoán theo NỘI DUNG — cột mã dài nhất = SKU, cột mã đầy đủ còn lại = model, còn lại = biến thể.
      * Cột có mã thuộc tính (vd deep_tskt_master, utilities_filter_master) -> spec PIM cũ để đối chiếu (TỰ LỌC TSKT).
    -> {"import": DataFrame, "spec": DataFrame, "ghi_chu": [..], "nhan_cot": {trường: tên cột}}"""
    rows = [list(r or []) for r in rows if r and any(v not in (None, "") for v in r)]
    out = {"import": pd.DataFrame(columns=COT_IMPORT), "spec": pd.DataFrame(columns=COT_SPEC), "ghi_chu": [],
           "nhan_cot": {}}
    if not rows:
        out["ghi_chu"].append("Không có dữ liệu.")
        return out
    n = max(len(r) for r in rows)
    rows = [r + [None] * (n - len(r)) for r in rows]
    vt, hang_tieu_de = {}, -1
    for i in range(min(3, len(rows))):
        j = tim_cot_truong([ep_text(x) for x in rows[i]], "sku")
        if j >= 0:
            hang_tieu_de = i
            vt = {f: tim_cot_truong([ep_text(x) for x in rows[i]], f) for f in ("sku", "model_code", "variant_code",
                                                                                 "category_code", "family_variant_code")}
            break
    if hang_tieu_de >= 0:
        h1 = [chuan_hoa_key(x) for x in rows[hang_tieu_de]]
        bd = hang_tieu_de + 1
        h2 = [""] * n
        if bd < len(rows) and _la_dong_ten_vn(rows[bd], vt["sku"]):
            h2 = [chuan_hoa_key(x) for x in rows[bd]]
            bd += 1
        data = rows[bd:]
    else:
        h1, h2, data = [""] * n, [""] * n, rows
        thong_ke = []
        for j in range(n):
            vs = [chuan_hoa_code(r[j]) for r in data]
            co = [v for v in vs if v]
            if not co or sum(bool(_MA_RE.match(v)) for v in co) < 0.8 * len(co):
                continue
            thong_ke.append((j, len(co) / len(data), sum(map(len, co)) / len(co),
                             sum(any(ch.isalpha() for ch in v) for v in co) / len(co)))
        if not thong_ke:
            out["ghi_chu"].append("Không nhận ra cột mã SKU nào (cần cột chỉ gồm chữ/số không dấu).")
            return out
        so = [t for t in thong_ke if t[3] < 0.5] or thong_ke     # ưu tiên cột toàn số làm SKU/model
        sku = max(so, key=lambda t: (t[1] >= 0.9, t[2]))
        con = [t for t in thong_ke if t[0] != sku[0]]
        vt = {"sku": sku[0], "model_code": -1, "variant_code": -1, "category_code": -1}
        if con:
            m = max(con, key=lambda t: (t[3] < 0.5, t[1]))
            vt["model_code"] = m[0]
            con = [t for t in con if t[0] != m[0]]
        if con:
            vt["variant_code"] = con[0][0]
        ten = {f: chr(65 + j) for f, j in vt.items() if j >= 0}
        out["ghi_chu"].append("Không có dòng tiêu đề — tự nhận theo nội dung: " +
                              ", ".join(f"cột {c} = {f}" for f, c in ten.items()) + ". Kiểm tra lại bảng xem trước.")
    out["nhan_cot"] = {f: (h1[j] or chr(65 + j)) for f, j in vt.items() if j >= 0}
    co_dinh = {j for j in vt.values() if j >= 0}
    n_mau = 0
    cot_spec = [(j, h1[j], h2[j]) for j in range(n) if j not in co_dinh and h1[j] and _MA_COT_RE.match(h1[j])
                and h1[j].lower() not in COT_KHONG_PHAI_SPEC]
    imp, spec, seen, trung = [], [], set(), 0
    for r in data:
        def g(f):
            j = vt.get(f, -1)
            return chuan_hoa_key(r[j]) if j >= 0 else ""
        sku = chuan_hoa_code(g("sku"))
        if not sku or not _MA_RE.match(sku):
            continue
        if sku in seen:
            trung += 1
            continue
        seen.add(sku)
        model, variant, cate = g("model_code"), g("variant_code"), chuan_hoa_id(g("category_code"))
        if variant and la_bien_the_mau(g("family_variant_code")):
            variant, n_mau = "", n_mau + 1
        imp.append([model, sku, variant, cate])
        for j, code, ten in cot_spec:
            v = lam_sach_gia_tri_pim(r[j], code)
            if v:
                spec.append([sku, model, cate, variant, code, ten, v])
    out["import"] = pd.DataFrame(imp, columns=COT_IMPORT)
    out["spec"] = pd.DataFrame(spec, columns=COT_SPEC)
    if trung:
        out["ghi_chu"].append(f"Bỏ {trung} dòng trùng SKU (giữ dòng đầu).")
    if n_mau:
        out["bien_the_mau"] = n_mau
        out["ghi_chu"].append(f"{n_mau} SKU có Mã họ biến thể chứa 'color' (biến thể theo màu) -> xuất vào file MODEL, không phải BIENTHE.")
    if cot_spec:
        out["ghi_chu"].append(f"Có {len(cot_spec)} cột thông số (TSKT/FILTER) -> tách thành spec PIM cũ để đối chiếu "
                              f"({len(spec):,} ô có giá trị).")
    if not imp:
        out["ghi_chu"].append("Không có SKU hợp lệ nào.")
    return out


def doc_ds_sku(rows: List[list]) -> Tuple[pd.DataFrame, Optional[str]]:
    r = doc_mot_cuc(rows)
    return r["import"], (" ".join(r["ghi_chu"]) if not len(r["import"]) else None)


def doc_import_don_gian(data: bytes, ten_file: str) -> Tuple[pd.DataFrame, Optional[str]]:
    """Danh sách SKU (tên cột tuỳ ý / không tiêu đề) -> IMPORT."""
    return doc_ds_sku(_doc_bang_tho(data, ten_file))


def doc_text_dan(txt: str) -> List[list]:
    """Chữ dán từ Excel (tab) / CSV (; ,) -> các dòng."""
    dong = [x for x in (txt or "").splitlines() if x.strip()]
    if not dong:
        return []
    sep = "\t" if any("\t" in x for x in dong) else (";" if any(";" in x for x in dong) else
                                                     ("," if any("," in x for x in dong) else None))
    return [[c.strip() for c in (x.split(sep) if sep else [x])] for x in dong]


def doc_cau_hinh_ngang(rows: List[list]) -> Dict[str, dict]:
    """Sheet CẤU HÌNH CATEGORY bố cục ngang của bản desktop: dòng có mã ngành ở
    cột A + tên cột B + các mã TSKT từ cột C; dòng ngay dưới (cột A trống) là
    tên tiếng Việt tương ứng."""
    cfg: Dict[str, dict] = {}
    cate = None
    cot_ma: Optional[Dict[int, str]] = None
    for r in rows[1:]:
        r = list(r or [])
        rid = chuan_hoa_id(r[0] if r else "")
        name = chuan_hoa_key(r[1] if len(r) > 1 else "")
        if rid and rid != cate:
            cate, cot_ma = rid, None
        if not cate:
            continue
        o = cfg.setdefault(cate, {"ten": name, "cot": [], "ten_cot": {}})
        if name and not o["ten"]:
            o["ten"] = name
        cells = [(c, chuan_hoa_key(r[c])) for c in range(2, len(r)) if chuan_hoa_key(r[c])]
        if not cells:
            continue
        la_dong_ma = all(re.fullmatch(r"[A-Za-z0-9_]+", v) for _, v in cells) and any("_" in v for _, v in cells)
        if la_dong_ma:
            cot_ma = {}
            for c, v in cells:
                cot_ma[c] = v
                if v not in o["cot"]:
                    o["cot"].append(v)
        elif cot_ma:
            for c, v in cells:
                code = cot_ma.get(c)
                if code and code not in o["ten_cot"]:
                    o["ten_cot"][code] = v
    return cfg


def doc_cau_hinh_columns_text(ds: List[dict]) -> Dict[str, dict]:
    """Định dạng [{cate_id, cate_name, columns_text: 'ma:Tên; ma:Tên'}]."""
    cfg = {}
    for x in ds:
        cate = chuan_hoa_id(x.get("cate_id"))
        if not cate:
            continue
        o = cfg.setdefault(cate, {"ten": chuan_hoa_key(x.get("cate_name")), "cot": [], "ten_cot": {}})
        for phan in str(x.get("columns_text", "")).split(";"):
            if ":" not in phan:
                continue
            ma, ten = phan.split(":", 1)
            ma = chuan_hoa_key(ma)
            if not ma or ma.lower() in ("model_activated", "variant_activated"):
                continue
            if ma not in o["cot"]:
                o["cot"].append(ma)
            o["ten_cot"].setdefault(ma, chuan_hoa_key(ten))
    return cfg


def cau_hinh_tu_template_pim(data: bytes, ten_file: str, cate: str, ten_cate: str = "") -> Dict[str, dict]:
    """Lấy danh sách cột của 1 ngành hàng từ chính file template/export PIM."""
    rows = _doc_bang_tho(data, ten_file)
    h1 = [chuan_hoa_key(h) for h in rows[0]] if rows else []
    h2 = ([chuan_hoa_key(h) for h in rows[1]] if len(rows) > 1 else []) + [""] * len(h1)
    cot = [h for h in h1 if h and h.lower() not in COT_KHONG_PHAI_SPEC]
    return {chuan_hoa_id(cate): {"ten": ten_cate, "cot": cot,
                                 "ten_cot": {h: h2[j] for j, h in enumerate(h1) if h in cot}}}


def doc_mapping_tskt(rows: List[list]) -> Tuple[pd.DataFrame, Optional[str]]:
    header = [ep_text(h) for h in rows[0]] if rows else []
    c = {"cate": _tim_cot(header, "Mã ngành hàng CMS", "Mã ngành hàng"),
         "cate_name": _tim_cot(header, "Tên ngành hàng CMS", "Tên ngành hàng"),
         # "MÃ THUỘC TÍNH TSKT" / "TÊN MASTER": tiêu đề của file mẫu mới (TEST HÀNG LOẠT IMPORT THÔNG SỐ NEW)
         "prop_id": _tim_cot(header, "Mã thuộc tính", "Mã thuộc tính TSKT"),
         "prop_name": _tim_cot(header, "Tên thuộc tính TSKT", "Tên thuộc tính"),
         "ma": _tim_cot(header, "Mã TSKT (MASTER)", "Mã MASTER"),
         "ten_ma": _tim_cot(header, "Tên TSKT (MASTER)", "Tên MASTER")}
    if min(c["cate"], c["prop_id"], c["ma"]) < 0:
        if len(header) < 5:
            return pd.DataFrame(columns=COT_MAP_TSKT), ("Thiếu cột bắt buộc: Mã ngành hàng CMS / Mã thuộc tính / "
                                                        "Mã TSKT (MASTER)")
        # tên cột khác -> đọc THEO VỊ TRÍ như bản desktop: A mã ngành, B tên ngành, C mã thuộc tính, D tên, E mã PIM, F tên
        c = {"cate": 0, "cate_name": 1, "prop_id": 2, "prop_name": 3, "ma": 4, "ten_ma": 5}
    out = []
    for r in rows[1:]:
        def g(k):
            j = c[k]
            return r[j] if 0 <= j < len(r) else ""
        cate, pid, ma = chuan_hoa_id(g("cate")), chuan_hoa_id(g("prop_id")), chuan_hoa_key(g("ma"))
        if cate and pid and ma:
            out.append([cate, chuan_hoa_key(g("cate_name")), pid, chuan_hoa_key(g("prop_name")), ma,
                        chuan_hoa_key(g("ten_ma"))])
    return pd.DataFrame(out, columns=COT_MAP_TSKT), None


def doc_ma_ho(rows: List[list]) -> Dict[str, str]:
    """Cột "MÃ HỌ" (family) của MAPPING TSKT MOI mẫu mới -> {mã ngành: mã họ} (giá trị đầu tiên)."""
    header = [ep_text(h) for h in rows[0]] if rows else []
    c_cate, c_ho = _tim_cot(header, "Mã ngành hàng CMS", "Mã ngành hàng"), _tim_cot(header, "Mã họ", "family_code")
    out: Dict[str, str] = {}
    if c_cate < 0 or c_ho < 0:
        return out
    for r in rows[1:]:
        if len(r) > max(c_cate, c_ho):
            cate, ho = chuan_hoa_id(r[c_cate]), chuan_hoa_key(r[c_ho])
            if cate and ho:
                out.setdefault(cate, ho)
    return out


def tim_cate_tab(ten_tab: str, cot: List[str], cate_map: set, ma_theo_cate: Dict[str, set]) -> Optional[str]:
    """Y hệt tim_tab_tskt() desktop: tab "TSKT <mã ngành> …" -> mã ngành (nếu có trong mapping);
    không thì ngành có NHIỀU mã cột trùng nhất."""
    phan = chuan_hoa_key(ten_tab[4:]).strip() if ten_tab.lower().startswith("tskt ") else ""
    tk = chuan_hoa_id(phan.split()[0]) if phan else ""
    if tk and tk in cate_map:
        return tk
    best, cate = 0, None
    for c, ma_set in ma_theo_cate.items():
        d = sum(1 for m in cot if m in ma_set)
        if d > best:
            best, cate = d, c
    return cate


def doc_mapping_filter(rows: List[list]) -> Tuple[pd.DataFrame, Optional[str]]:
    header = [ep_text(h) for h in rows[0]] if rows else []
    c = {"cate": _tim_cot(header, "MÃ NGÀNH HÀNG CMS", "Mã ngành hàng"),
         "cate_name": _tim_cot(header, "TÊN NGÀNH HÀNG CMS", "Tên ngành hàng"),
         "prop_id": _tim_cot(header, "MÃ THUỘC TÍNH FILTER", "Mã thuộc tính"),
         "prop_name": _tim_cot(header, "TÊN THUỘC TÍNH CMS", "Tên thuộc tính FILTER", "Tên thuộc tính"),
         "moi": _tim_cot(header, "Mã thuộc tính mới"), "cu": _tim_cot(header, "Mã thuộc tính cũ")}
    if min(c["cate"], c["prop_id"]) < 0 or (c["moi"] < 0 and c["cu"] < 0):
        if len(header) < 5:
            return pd.DataFrame(columns=COT_MAP_FILTER), ("Thiếu cột bắt buộc: MÃ NGÀNH HÀNG CMS / MÃ THUỘC TÍNH "
                                                          "FILTER / Mã thuộc tính mới (hoặc cũ)")
        # theo VỊ TRÍ như desktop: A mã ngành, B tên, C mã thuộc tính filter, D tên, E mã cũ, F mã mới
        c = {"cate": 0, "cate_name": 1, "prop_id": 2, "prop_name": 3, "cu": 4, "moi": 5 if len(header) > 5 else 4}
    out = []
    for r in rows[1:]:
        def g(k):
            j = c[k]
            return r[j] if 0 <= j < len(r) else ""
        cate, pid = chuan_hoa_id(g("cate")), chuan_hoa_id(g("prop_id"))
        ma = chuan_hoa_key(g("moi")) or chuan_hoa_key(g("cu"))
        if cate and pid and ma:
            out.append([cate, chuan_hoa_key(g("cate_name")), pid, chuan_hoa_key(g("prop_name")), ma])
    return pd.DataFrame(out, columns=COT_MAP_FILTER), None


def doc_data_pim(rows: List[list]) -> Tuple[pd.DataFrame, Optional[str]]:
    """DATA PIM: cột A Code, B Name, C Type, D Group, E Active, F OptionCode, G OptionValue.
    Có/không có dòng tên tiếng Việt thứ 2 đều được."""
    out = []
    for r in rows[1:]:
        r = list(r or []) + [None] * 7
        code = chuan_hoa_key(r[0])
        if not code or code == "Mã thuộc tính":
            continue
        out.append([code, chuan_hoa_key(r[1]), chuan_hoa_key(r[2]), chuan_hoa_key(r[3]), ep_text(r[4]),
                    chuan_hoa_id(r[5]), chuan_hoa_key(r[6])])
    if not out:
        return pd.DataFrame(columns=COT_DATA_PIM), "Không có dòng dữ liệu."
    return pd.DataFrame(out, columns=COT_DATA_PIM), None


def doc_workspace_cu(data: bytes, ten_file: str = "du_lieu_pim.xlsx") -> dict:
    """CHUYỂN NHÀ 1 CHẠM: đọc toàn bộ workspace của bản desktop (du_lieu_pim.xlsx)."""
    so = SoExcel(data, ten_file)
    ten = so.sheets
    kq: dict = {"canh_bao": []}

    def rows_of(s):
        return so.rows(s) if s in ten else None
    r = rows_of("CẤU HÌNH CATEGORY")
    if r:
        kq["cau_hinh"] = doc_cau_hinh_ngang(r)
    r = rows_of("MAPPING TSKT MOI")
    if r:
        kq["map_tskt"], e = doc_mapping_tskt(r)
        e and kq["canh_bao"].append(f"MAPPING TSKT MOI: {e}")
    r = rows_of("MAPPING FILTER MOI")
    if r:
        kq["map_filter"], e = doc_mapping_filter(r)
        e and kq["canh_bao"].append(f"MAPPING FILTER MOI: {e}")
    r = rows_of("MAPPING TSKT MOI")
    if r:
        ho = doc_ma_ho(r)
        if ho:
            kq["ma_ho"] = ho
    # ---- Tab "TSKT <mã> <tên>" (bản desktop dùng tab làm danh sách cột, ưu tiên hơn cấu hình) ----
    bo_qua = {"import", "chọn ngành hàng", "xử lý hàm", "cấu hình category", "log", "data sp", "data pim",
              "thuộc tính cms/pim", "thuộc tính cmspim", "mapping tskt moi", "mapping filter moi", "import filter"}
    mt, mf = kq.get("map_tskt"), kq.get("map_filter")
    ma_theo_cate: Dict[str, set] = defaultdict(set)
    for df in (mt, mf):
        if df is not None and len(df):
            for c, m in df[["cate", "ma"]].itertuples(index=False):
                ma_theo_cate[c].add(m)
    tabs: Dict[str, dict] = {}
    for tn in ten:
        n = tn.lower().strip()
        if n in bo_qua or not (n == "tskt" or n.startswith("tskt ")):
            continue
        rr = rows_of(tn) or []
        if not rr:
            continue
        h1 = [chuan_hoa_key(x) for x in rr[0]]
        h2 = [chuan_hoa_key(x) for x in (rr[1] if len(rr) > 1 else [])] + [""] * len(h1)
        seen, cot, ten_cot = set(), [], {}
        for j, h in enumerate(h1):
            if h and h not in seen:
                seen.add(h)
                if h.lower() not in COT_KHONG_DIEN:
                    cot.append(h)
                    if h2[j]:
                        ten_cot[h] = h2[j]
        if not cot:
            continue
        cate = tim_cate_tab(tn, cot, set(ma_theo_cate), ma_theo_cate)
        if cate and cate not in tabs:
            tabs[cate] = {"tab": tn, "cot": cot, "ten_cot": ten_cot}
    if tabs:
        kq["tab_tskt"] = tabs
        ch = kq.setdefault("cau_hinh", {})
        for cate, t in tabs.items():
            cfg = ch.get(cate)
            if cfg and cfg.get("cot"):
                # desktop: cột tab + cột cấu hình còn thiếu (thêm cuối); cột chỉ có trên tab -> xuất trống, không điền
                cfg_cot = list(cfg["cot"])
                ngoai = [c for c in t["cot"] if c not in set(cfg_cot)]
                cot = t["cot"] + [c for c in cfg_cot if c not in set(t["cot"])]
                ch[cate] = {**cfg, "cot": cot, "ten_cot": {**t["ten_cot"], **cfg.get("ten_cot", {})},
                            "cot_khong_dien": ngoai, "nguon": f"tab {t['tab']} + cấu hình"}
            else:
                ten_nh = re.sub(r"^\s*TSKT\s+\S+\s*", "", t["tab"]).strip()
                ch[cate] = {"ten": ten_nh, "cot": t["cot"], "ten_cot": t["ten_cot"], "nguon": f"tab {t['tab']}"}
    if kq.get("ma_ho") and kq.get("cau_hinh"):
        for cate, ho in kq["ma_ho"].items():
            if cate in kq["cau_hinh"]:
                kq["cau_hinh"][cate].setdefault("ma_ho", ho)
    r = rows_of("CHỌN NGÀNH HÀNG")
    if r and len(r) > 1:
        h = [chuan_hoa_key(x).upper() for x in r[0]]
        if "CHỌN" in h and "MÃ NH" in h:
            kq["chon_nganh"] = {chuan_hoa_id(x[h.index("MÃ NH")]): ep_text(x[h.index("CHỌN")]).strip().upper()
                                in ("TRUE", "1", "X", "CÓ") for x in r[1:] if len(x) > h.index("MÃ NH")
                                and chuan_hoa_id(x[h.index("MÃ NH")])}
    r = rows_of("DATA PIM")
    if r:
        kq["data_pim"], e = doc_data_pim(r)
    r = rows_of("IMPORT")
    if r and r[0]:
        df, e = doc_ds_sku(r)
        kq["import"] = df
        e and kq["canh_bao"].append(f"IMPORT: {e}")
    r = rows_of("DATA SP")
    for tn in sorted(x for x in ten if re.fullmatch(r"DATA SP \(\d+\)", x)):  # phần tách khi > 1 triệu dòng
        r = (r or []) + (rows_of(tn) or [])[1:]
    if r:
        out = [[chuan_hoa_id(x[0]), chuan_hoa_code(x[1]), chuan_hoa_key(x[2]), chuan_hoa_id(x[3]),
                chuan_hoa_key(x[4]), chuan_hoa_key(x[5]), chuan_hoa_id(x[6] if len(x) > 6 else "")]
               for x in r[1:] if x and len(x) > 5 and chuan_hoa_code(x[1])]
        kq["data_sp"] = pd.DataFrame(out, columns=COT_DATA_SP)
    r = rows_of("SPEC PIM TẠM")
    if r:
        out = []
        for x in r[1:]:
            x = list(x or []) + [None] * 7
            if chuan_hoa_code(x[0]):
                code = chuan_hoa_key(x[4])
                out.append([chuan_hoa_code(x[0]), chuan_hoa_key(x[1]), chuan_hoa_id(x[2]), chuan_hoa_key(x[3]),
                            code, chuan_hoa_key(x[5]), lam_sach_gia_tri_pim(x[6], code)])
        kq["spec"] = pd.DataFrame(out, columns=COT_SPEC)
    r = rows_of("ĐƠN VỊ")
    if r:
        dv = {}
        for x in r[1:]:
            x = list(x or []) + [None] * 5
            if chuan_hoa_id(x[0]) and chuan_hoa_key(x[2]) and chuan_hoa_key(x[4]):
                dv[f"{chuan_hoa_id(x[0])}\t{chuan_hoa_key(x[2])}"] = chuan_hoa_key(x[4])
        kq["don_vi"] = dv
    kq["sheets"] = ten
    return kq


# ============================================================================
# §4 THAM CHIẾU (data/reference) — map theo TÊN + tách kích thước ghép
# ============================================================================
_REF_CACHE: dict = {}


def _ref(ten: str):
    if ten not in _REF_CACHE:
        f = REF_DIR / ten
        try:
            _REF_CACHE[ten] = json.loads(f.read_text(encoding="utf-8")) if f.exists() else None
        except Exception:  # noqa: BLE001
            _REF_CACHE[ten] = None
    return _REF_CACHE[ten]


def map_theo_ten() -> Dict[str, Dict[str, str]]:
    """{cate: {tên thuộc tính CMS (chuẩn hoá): mã master}} — gộp
    thuoc_tinh_theo_ten.json + ten_cms_pim.json (chỉ dòng có mã ngành số)."""
    if "_map_ten" in _REF_CACHE:
        return _REF_CACHE["_map_ten"]
    out: Dict[str, Dict[str, str]] = defaultdict(dict)
    for cate, v in (_ref("thuoc_tinh_theo_ten.json") or {}).items():
        for a in v.get("attrs", []):
            for ten in (a.get("tskt_name"), a.get("master_name")):
                if ten and a.get("master_code"):
                    out[chuan_hoa_id(cate)].setdefault(chuan_hoa_ten(ten), a["master_code"])
    for x in _ref("ten_cms_pim.json") or []:
        cate, ma = chuan_hoa_id(x.get("cate_id")), chuan_hoa_key(x.get("pim_code"))
        if re.fullmatch(r"\d+", cate or "") and ma and ma.lower() != "done" and "_" in ma:
            out[cate].setdefault(chuan_hoa_ten(x.get("cms_attr_name")), ma)
    _REF_CACHE["_map_ten"] = dict(out)
    return _REF_CACHE["_map_ten"]


def cau_hinh_tham_chieu() -> Dict[str, dict]:
    return doc_cau_hinh_columns_text(_ref("cau_hinh_cot_nganh_hang.json") or [])


def nhan_kich_thuoc_ghep(cate: str) -> List[str]:
    """Các nhãn con (Cao/Ngang/Sâu/Nặng...) của thuộc tính ghép kích thước theo ngành hàng."""
    nhan: List[str] = []
    for x in (_ref("kich_thuoc_ghep.json") or {}).get(chuan_hoa_id(cate), []):
        for s in x.get("sub_labels", []):
            if s not in nhan:
                nhan.append(s)
    for x in (_ref("kich_thuoc_ghep_alias.json") or {}).get(chuan_hoa_id(cate), []):
        if "kích thước" in chuan_hoa_ten(x.get("container", "")) or "khối lượng" in chuan_hoa_ten(
                x.get("container", "")):
            for s in x.get("sub_labels", []):
                if s not in nhan:
                    nhan.append(s)
    return nhan or ["Ngang", "Cao", "Sâu", "Dài", "Rộng", "Dày", "Dọc", "Nặng", "Khối lượng"]


_SO_DV_RE = r"(\d+(?:[.,]\d+)?)\s*(cm|mm|m|kg|g|inch|\")?"


def tach_kich_thuoc_ghep(text: str, nhan: List[str]) -> Dict[str, str]:
    """'Ngang 122.6 cm - Cao 71.1 cm - Dày 7.2 cm' -> {'Ngang': '122.6 cm', ...}.
    Chỉ lấy nhãn xuất hiện THẬT trong chuỗi (không suy đoán theo vị trí)."""
    s = chuan_hoa_key(text)
    out: Dict[str, str] = {}
    for n in sorted(nhan, key=len, reverse=True):
        m = re.search(rf"(?<![\wÀ-ỹ]){re.escape(n)}\s*:?\s*{_SO_DV_RE}", s, re.IGNORECASE)
        if m and n not in out:
            out[n] = (m.group(1) + (" " + m.group(2) if m.group(2) else "")).strip()
    return out


# ============================================================================
# §5 MAP DỮ LIỆU (thay ① Điền dữ liệu)
# ============================================================================
def option_maps(data_pim: pd.DataFrame) -> dict:
    """option_map {(mã filter, tên thường): mã option}, opt_ten {(mã filter, mã option): tên},
    opt_ds {mã filter: [(mã option, tên)]}, ten_filter {mã: tên}."""
    om: Dict[Tuple[str, str], str] = {}
    ot: Dict[Tuple[str, str], str] = {}
    ods: Dict[str, List[Tuple[str, str]]] = defaultdict(list)
    tf: Dict[str, str] = {}
    if data_pim is not None and len(data_pim):
        for code, name, oc, ov in data_pim[["Code", "Name", "OptionCode", "OptionValue"]].itertuples(index=False):
            if not code or not ov:
                continue
            om.setdefault((code, khoa_option(ov)), oc)
            if oc and (code, oc) not in ot:
                ot[(code, oc)] = ov
                ods[code].append((oc, ov))
            tf.setdefault(code, name)
    return {"option_map": om, "opt_ten": ot, "opt_ds": dict(ods), "ten_filter": tf}


def chay_map(data_sp: pd.DataFrame, imp: pd.DataFrame, cau_hinh: Dict[str, dict],
             map_tskt: pd.DataFrame, map_filter: pd.DataFrame, opt: dict,
             map_ten: str = "tat", quy_doi: Optional[Dict[str, str]] = None,
             sua_gt: Optional[Dict[str, str]] = None, sua_sku: Optional[Dict[str, dict]] = None) -> dict:
    """map_ten (dự phòng khi PROPERTYID chưa có trong MAPPING):
         "tat"        — không map theo tên (giống hệt bản desktop)
         "cau_hinh"   — chỉ khớp ĐÚNG tên cột của chính ngành hàng (CẤU HÌNH CATEGORY) — chính xác cao
         "tham_chieu" — thêm cả bảng tham chiếu data/reference (rộng hơn, cần soát lại)
    sua_gt  — quy tắc sửa giá trị CMS sai theo GIÁ TRỊ: {"cate\tma\tkhoa(giá trị cũ)": giá trị đúng}
    sua_sku — quy tắc sửa theo SKU: {"sku\tma": {"cu": khoa(giá trị cũ), "moi": giá trị đúng}};
              chỉ áp khi ô vẫn còn đúng giá trị cũ (CMS đã sửa thì tự bỏ qua).
    -> {"bang": {cate: bang}, "log": [...], "tom_tat": {...}, "chua_map": DataFrame}."""
    log: List[list] = []
    sku_rule: Dict[str, List[Tuple[str, dict]]] = defaultdict(list)
    for k, v in (sua_sku or {}).items():
        if k.count("\t") == 1 and isinstance(v, dict):
            a, b = k.split("\t")
            sku_rule[a].append((b, v))
    goc_cms: Dict[str, str] = {}  # "sku\tma" -> giá trị CMS gốc của ô đã sửa theo quy tắc
    gt_cot: Dict[str, set] = defaultdict(set)
    for k in (sua_gt or {}):
        if k.count("\t") == 2:
            a, b, _ = k.split("\t")
            gt_cot[a].add(b)
    imp = imp.drop_duplicates("sku")
    skus = list(imp.sku)
    sku_set = set(skus)
    sp = data_sp[data_sp.PRODUCTCODE.isin(sku_set)] if len(data_sp) else data_sp
    # --- mapping theo cate
    tm: Dict[str, Dict[str, str]] = defaultdict(dict)
    for cate, pid, ma in map_tskt[["cate", "prop_id", "ma"]].itertuples(index=False):
        tm[cate].setdefault(pid, ma)
    fm: Dict[str, Dict[str, str]] = defaultdict(dict)
    for cate, pid, ma in map_filter[["cate", "prop_id", "ma"]].itertuples(index=False):
        fm[cate].setdefault(pid, ma)
    cate_cua_pid: Dict[str, set] = defaultdict(set)
    for d in (tm, fm):
        for cate, m in d.items():
            for pid in m:
                cate_cua_pid[pid].add(cate)
    # --- phân loại SKU -> cate
    cate_khai: Dict[str, str] = {}
    pid_sku: Dict[str, set] = defaultdict(set)
    for sku, pid, cid in sp[["PRODUCTCODE", "PROPERTYID", "CATEGORYID"]].itertuples(index=False):
        if cid:
            if sku in cate_khai and cate_khai[sku] != cid:
                log.append([sku, f"{cate_khai[sku]} | {cid}", "", "CATEGORYID mâu thuẫn trong DATA SP — dùng giá trị đầu"])
            else:
                cate_khai.setdefault(sku, cid)
        if pid:
            pid_sku[sku].add(pid)
    theo_cate: Dict[str, List[str]] = defaultdict(list)
    nguon_cate: Dict[str, Counter] = defaultdict(Counter)  # cho bảng CHỌN NGÀNH HÀNG (như desktop)
    khong_data: List[str] = []
    co_map = set(tm) | set(fm)
    for sku in skus:
        c = cate_khai.get(sku)
        tu_doan = False
        if c and c not in co_map:  # như desktop phan_loai(): ngành không có trong mapping -> bỏ qua SKU
            khong_data.append(sku)
            log.append([sku, c, "", "CATEGORYID không tồn tại trong mapping — bỏ qua SKU"])
            continue
        if not c:
            diem = Counter(cc for pid in pid_sku.get(sku, ()) for cc in cate_cua_pid.get(pid, ()))
            if diem:
                xep = sorted(diem.items(), key=lambda kv: (-kv[1], kv[0]))
                c = xep[0][0]
                tu_doan = True
                if len(xep) > 1 and xep[1][1] == xep[0][1]:
                    nguon_cate[c]["tie"] += 1
                log.append([sku, c, "", "Không có CATEGORYID — tự nhận diện theo thuộc tính"])
        if not c:
            khong_data.append(sku)
            log.append([sku, "", "", "Không có dòng nào trong DATA SP"])
            continue
        theo_cate[c].append(sku)
        nguon_cate[c]["tu_doan" if tu_doan else "tu_cate_id"] += 1
    nl = {r.sku: r for r in imp.itertuples(index=False)}
    ten_map = map_theo_ten() if map_ten == "tham_chieu" else {}
    bang: Dict[str, dict] = {}
    chua_map: Dict[Tuple[str, str], dict] = {}
    st = Counter()
    for cate, ds in theo_cate.items():
        cfg = cau_hinh.get(cate)
        tmc, fmc, tenc = tm.get(cate, {}), fm.get(cate, {}), ten_map.get(cate, {})
        if not (cfg and cfg.get("cot")):  # như desktop: không có tab TSKT lẫn cấu hình -> bỏ qua ngành
            log.append(["", cate, "", "Không có tab TSKT lẫn cấu hình — bỏ qua"])
            khong_data.extend(ds)
            continue
        # cột = model_code, sku, variant_code + đúng thứ tự CẤU HÌNH (giữ cả family_code/category_code… như desktop,
        # để trống); chỉ điền dữ liệu vào cột thuộc tính (loại model_code/sku/category_code/variant_code)
        cot = [c for c in cfg["cot"] if c.lower() not in COT_CO_DINH_LOWER]
        ten_cot = dict(cfg.get("ten_cot", {}))
        for _, r in map_tskt[map_tskt.cate == cate][["ma", "ten_ma"]].iterrows():
            if r.ten_ma:
                ten_cot.setdefault(r.ma, r.ten_ma)
        # cột chỉ có trên tab desktop mà không có trong cấu hình gốc -> xuất trống, không điền (như desktop)
        khong_dien = set(cfg.get("cot_khong_dien") or ())
        for c in khong_dien:
            log.append(["", cate, c, "Cột có trên tab nhưng KHÔNG có trong cấu hình category — bỏ qua, không điền"])
        cot_set = {c for c in cot if c.lower() not in COT_KHONG_DIEN and c not in khong_dien}
        dsset = set(ds)
        bt: Dict[Tuple[str, str], List[str]] = defaultdict(list)
        bf: Dict[Tuple[str, str], List[str]] = defaultdict(list)
        theo_id: set = set()
        da_bao: set = set()
        sp_c = sp[sp.PRODUCTCODE.isin(dsset)]
        cho_ten: List[tuple] = []
        ten_cau_hinh: Dict[str, str] = {}
        if map_ten in ("cau_hinh", "tham_chieu"):
            dem_ten = Counter(chuan_hoa_ten(t) for t in ten_cot.values() if t)
            ten_cau_hinh = {chuan_hoa_ten(t): m for m, t in ten_cot.items()
                            if t and m in cot_set and dem_ten[chuan_hoa_ten(t)] == 1}
        for sku, pid, pname, val in sp_c[["PRODUCTCODE", "PROPERTYID", "PROPERTYNAME", "PROPVALUE"]].itertuples(
                index=False):
            if not pid or not val:
                continue
            da = False
            if pid in tmc:
                da = True
                ma = tmc[pid]
                if ma in cot_set:
                    theo_id.add((sku, ma))
                    if val not in bt[(sku, ma)]:
                        bt[(sku, ma)].append(val)
            if pid in fmc:
                da = True
                ma = fmc[pid]
                if ma in cot_set:
                    oc = (quy_doi or {}).get(f"{ma}\t{khoa_quy_doi(val)}") or opt["option_map"].get((ma, khoa_option(val)))
                    if oc:
                        if oc not in bf[(sku, ma)]:
                            bf[(sku, ma)].append(oc)
                    elif (ma, val) not in da_bao:
                        da_bao.add((ma, val))
                        st["filter_khong_khop"] += 1
                        log.append([sku, cate, ma, f'FILTER: giá trị CMS "{val}" chưa có option trong DATA PIM — để trống'])
            if not da:
                cho_ten.append((sku, pid, pname, val))
        # --- dự phòng: map theo TÊN thuộc tính (chỉ vào cột ĐÃ cấu hình, không ghi đè ô map theo mã)
        for sku, pid, pname, val in cho_ten:
            ma = ten_cau_hinh.get(chuan_hoa_ten(pname)) or (tenc.get(chuan_hoa_ten(pname)) if tenc else None)
            if ma and ma in cot_set and not la_cot_filter(ma) and (sku, ma) not in theo_id:
                if val not in bt[(sku, ma)]:
                    bt[(sku, ma)].append(val)
                    st["map_theo_ten"] += 1
                k = ("ten", cate, pid)
                if k not in chua_map:
                    chua_map[k] = {"cate": cate, "prop_id": pid, "prop_name": pname, "so_dong": 0, "vi_du": val,
                                   "goi_y_ma": ma, "trang_thai": "Đã map theo TÊN (nên thêm vào MAPPING TSKT)"}
                chua_map[k]["so_dong"] += 1
            else:
                k = ("chua", cate, pid)
                if k not in chua_map:
                    chua_map[k] = {"cate": cate, "prop_id": pid, "prop_name": pname, "so_dong": 0, "vi_du": val,
                                   "goi_y_ma": ma or "", "trang_thai": "CHƯA MAP — bỏ qua"}
                chua_map[k]["so_dong"] += 1
        rows = []
        for i, sku in enumerate(ds):
            r = nl.get(sku)
            vals = {"model_code": r.model_code if r else "", "sku": sku, "variant_code": r.variant_code if r else ""}
            for ma in cot:
                if bt.get((sku, ma)):
                    vals[ma] = SEP_TSKT.join(bt[(sku, ma)])
                elif bf.get((sku, ma)):
                    vals[ma] = SEP_FILTER.join(bf[(sku, ma)])
                else:
                    vals[ma] = ""
            for ma, rule in sku_rule.get(sku, ()):
                if ma not in cot_set:
                    continue
                if khoa_quy_doi(vals.get(ma, "")) == rule.get("cu", ""):
                    goc_cms[f"{sku}\t{ma}"] = vals.get(ma, "")
                    vals[ma] = rule.get("moi", "")
                    st["sua_duyet"] += 1
                else:
                    log.append([sku, cate, ma, "Quy tắc sửa theo SKU KHÔNG áp: giá trị CMS đã khác lúc đề xuất"])
            for ma in gt_cot.get(cate, ()):
                v0 = vals.get(ma, "")
                if v0 and ma in cot_set and f"{sku}\t{ma}" not in goc_cms:
                    m = sua_gt.get(f"{cate}\t{ma}\t{khoa_quy_doi(v0)}")
                    if m is not None and m != v0:
                        goc_cms.setdefault(f"{sku}\t{ma}", v0)
                        vals[ma] = m
                        st["sua_duyet"] += 1
            rows.append({"sku": sku, "model": vals["model_code"], "variant": vals["variant_code"],
                         "cate_pim": r.category_code if r else "", "so_dong": i + 3, "vals": vals})
        ten_nh = (cfg or {}).get("ten") or next(
            (x for x in map_tskt[map_tskt.cate == cate].cate_name if x), "")
        bang[cate] = {"title": f"TSKT {cate} {ten_nh}".strip(), "cate": cate, "ten_nh": ten_nh, "attr": cot,
                      "ten": ten_cot, "rows": rows}
        st["so_sku"] += len(ds)
        st["so_o"] += len(bt) + len(bf)
    df_cm = pd.DataFrame(list(chua_map.values()),
                         columns=["cate", "prop_id", "prop_name", "so_dong", "vi_du", "goi_y_ma", "trang_thai"])
    chon = []
    for c, ds_ in sorted(theo_cate.items(), key=lambda kv: -len(kv[1])):
        n = nguon_cate[c]
        cfg = cau_hinh.get(c) or {}
        ghi = []
        if c not in bang:
            ghi.append("Chưa có tab TSKT lẫn cấu hình — thêm 1 trong 2")
        if n["tie"]:
            ghi.append("Có SKU tự nhận diện trùng điểm (xem LOG)")
        chon.append({"CHỌN": c in bang, "MÃ NH": c, "TÊN NGÀNH HÀNG": cfg.get("ten") or (bang.get(c) or {}).get("ten_nh", ""),
                     "SỐ SKU": len(ds_), "NGUỒN CATE": ("CATEGORYID + tự nhận diện" if n["tu_cate_id"] and n["tu_doan"]
                                                       else "CATEGORYID" if n["tu_cate_id"] else "tự nhận diện"),
                     "TAB TSKT / CẤU HÌNH": cfg.get("nguon") or (f"cấu hình — {len(cfg.get('cot', []))} cột" if cfg else ""),
                     "GHI CHÚ": " | ".join(ghi)})
    return {"bang": bang, "log": log, "khong_data": khong_data, "chua_map": df_cm, "goc_cms": goc_cms, "chon": chon,
            "tom_tat": {"so_sku": st["so_sku"], "so_nganh": len(bang), "so_o": st["so_o"],
                        "map_theo_ten": st["map_theo_ten"], "filter_khong_khop": st["filter_khong_khop"],
                        "khong_data": len(khong_data), "o_sua_theo_quy_tac": st["sua_duyet"]}, "luc": bay_gio()}


@lru_cache(maxsize=_CACHE)
def _filter_co_chu(moi: str) -> bool:
    """Ô FILTER có phần tử KHÔNG phải mã số (PIM chỉ nhận mã option)."""
    return any(not re.fullmatch(r"\d+", chuan_hoa_id(x.strip())) for x in re.split(r"[,|]", moi) if x.strip())


@lru_cache(maxsize=_CACHE)
def _gon(s: str) -> str:
    """Gộp khoảng trắng (so khớp chuỗi con khi đối soát)."""
    return _WS_RE.sub(" ", s)


def cot_tt(b: dict) -> List[str]:
    """Cột THUỘC TÍNH của 1 bảng (bỏ family_code/category_code/*_activated… — vẫn xuất trống như desktop
    nhưng không kiểm tra/đối soát)."""
    return [m for m in b["attr"] if m.lower() not in COT_KHONG_PHAI_SPEC]


def model_trung(bang: Dict[str, dict]) -> List[Tuple[str, str, int]]:
    """model_code xuất hiện ở NHIỀU dòng file MODEL (nhiều SKU không biến thể cùng 1 model) — giống desktop vẫn
    xuất đủ các dòng, nhưng PIM sẽ lấy dòng cuối cùng -> cảnh báo để người dùng biết. -> [(cate, model, số dòng)]."""
    out = []
    for cate, b in bang.items():
        dem = Counter(r["model"] for r in b["rows"] if r["model"] and not chuan_hoa_key(r["vals"].get("variant_code")))
        out += [(cate, m, n) for m, n in dem.items() if n > 1]
    return out


def ten_tab_desktop(title: str) -> str:
    """Y hệt _safe_sheet_title() bản desktop (tên tab = tên file xuất): bỏ ký tự cấm, cắt 31 ký tự.
    Thêm thay "<>| để tên file hợp lệ trên Windows."""
    t = re.sub(r'[\\/*?:\[\]]', "_", (title or "")[:90]).strip()
    t = re.sub(r'["<>|]', "_", t)
    return t[:31] if t else "Sheet"


def nen_df(df: pd.DataFrame, nguong: float = 0.5) -> pd.DataFrame:
    """Giảm RAM cho bảng lớn (DATA SP, spec): cột chữ lặp nhiều -> kiểu category (giá trị y nguyên khi đọc).
    DATA SP 2,4 triệu dòng: ~1,1 GB -> ~0,2 GB. CHỈ dùng cho bảng không sửa từng ô."""
    if df is None or len(df) < 20000:
        return df
    out = df.copy()
    for c in out.columns:
        if out[c].dtype == object and out[c].nunique(dropna=False) <= nguong * len(out):
            out[c] = out[c].astype("category")
    return out


def bang_sang_df(bang: Dict[str, dict]) -> pd.DataFrame:
    """Lưu trữ kết quả map dạng dọc (cate, sku, model, variant, cate_pim, ma, gia_tri) + meta."""
    out = []
    for cate, b in bang.items():
        for r in b["rows"]:
            for ma in b["attr"]:
                out.append([cate, r["sku"], r["model"], r["variant"], r["cate_pim"], ma, r["vals"].get(ma, "")])
    return pd.DataFrame(out, columns=["cate", "sku", "model", "variant", "cate_pim", "ma", "gia_tri"])


def df_sang_bang(df: pd.DataFrame, meta: Dict[str, dict]) -> Dict[str, dict]:
    bang: Dict[str, dict] = {}
    if df is None or not len(df):
        return bang
    for cate, g in df.groupby("cate", sort=False):
        m = meta.get(cate, {})
        attr = m.get("attr") or list(dict.fromkeys(g.ma))
        rows_by: Dict[str, dict] = {}
        for sku, model, variant, cp, ma, v in g[["sku", "model", "variant", "cate_pim", "ma", "gia_tri"]].itertuples(
                index=False):
            r = rows_by.get(sku)
            if r is None:
                r = rows_by[sku] = {"sku": sku, "model": model, "variant": variant, "cate_pim": cp,
                                    "so_dong": len(rows_by) + 3,
                                    "vals": {"model_code": model, "sku": sku, "variant_code": variant}}
            r["vals"][ma] = v
        bang[cate] = {"title": m.get("title", f"TSKT {cate}"), "cate": cate, "ten_nh": m.get("ten_nh", ""),
                      "attr": attr, "ten": m.get("ten", {}), "rows": list(rows_by.values())}
    return bang


# ============================================================================
# §6 KIỂM TRA + ĐỐI CHIẾU (tính lại tức thì khi sửa)
# ============================================================================
def chuan_bi_spec(spec: pd.DataFrame) -> Tuple[Dict[str, dict], Dict[Tuple[str, str], dict]]:
    cu_sku: Dict[str, dict] = {}
    if spec is not None and len(spec):
        for sku, model, cate, variant, ma, v in spec[["sku", "model_code", "category_code", "variant_code", "ma",
                                                       "gia_tri"]].itertuples(index=False):
            e = cu_sku.get(sku)
            if e is None:
                e = cu_sku[sku] = {"sku": sku, "model": model, "cate": cate, "variant": variant, "spec": {}}
            if ma:
                e["spec"][ma] = lam_sach_gia_tri_pim(v, ma)
    cu_mv: Dict[Tuple[str, str], dict] = {}
    for e in cu_sku.values():
        if e["model"]:
            cu_mv.setdefault((e["model"], e["variant"]), e)
    return cu_sku, cu_mv


def bien_doi_o(cate: str, khoa: str, code: str, goc: str, sua: dict, don_vi: dict, rong: dict
               ) -> Tuple[str, Optional[str]]:
    """Giá trị cuối của 1 ô: sửa tay > quy tắc Không/Đang cập nhật > đơn vị (rule cũ) > biến đổi hàng loạt.
    -> (giá trị, loại biến đổi: 'sua'|'trong'|'thay'|'dv'|None)."""
    key = (cate, khoa, code)
    if key in sua:
        return chuan_hoa_key(sua[key]), "sua"
    moi, hd = ap_quy_tac_rong(goc, code, rong)
    if hd:
        return moi, hd
    if la_cot_filter(code):
        return goc, None
    v, loai = goc, None
    dv = don_vi.get((cate, code))
    if dv and isinstance(dv, str):
        m2 = them_don_vi(goc, dv)
        if m2 != goc:
            v, loai = m2, "dv"
    buoc = don_vi.get((cate, code, "bd")) or don_vi.get(("*", code, "bd"))
    if buoc:
        m3 = ap_bien_doi(v, buoc)
        if m3 != v:
            v, loai = m3, loai or "bd"
    return v, loai


def tinh_kiem_tra(bang: Dict[str, dict], imp: pd.DataFrame, spec: Optional[pd.DataFrame], opt: dict,
                  sua: Optional[dict] = None, don_vi: Optional[dict] = None, rong: Optional[dict] = None,
                  don_vi_luu: Optional[dict] = None) -> dict:
    sua, don_vi, rong, don_vi_luu = sua or {}, don_vi or {}, rong or {}, don_vi_luu or {}
    om = opt["option_map"]
    _tap_cache: Dict[Tuple[str, str], frozenset] = {}

    def tap(v: str, code: str) -> frozenset:
        k = (v, code if la_cot_filter(code) else "")
        x = _tap_cache.get(k)
        if x is None:
            x = _tap_cache[k] = tap_gia_tri(v, code, om)
        return x
    nl_sku = {r.sku: r for r in imp.itertuples(index=False)} if imp is not None and len(imp) else {}
    co_spec = spec is not None and len(spec) > 0
    cu_sku, cu_mv = chuan_bi_spec(spec) if co_spec else ({}, {})
    # tên cột dự phòng = dòng 2 file export PIM (như desktop: b["ten"].get(code) or ten_cu.get(code))
    ten_cu: Dict[str, str] = {}
    if co_spec:
        for m, t in spec[["ma", "ten"]].drop_duplicates("ma").itertuples(index=False):
            if m and t:
                ten_cu[m] = t
    khoa_rong_qt = _KHOA_RONG_MAC_DINH | set(rong)
    chi_tiet: Dict[str, dict] = {}
    khac: List[list] = []
    canh_bao: List[list] = []
    st: Counter = Counter()
    so_tron_cot: Dict[Tuple[str, str], List[str]] = defaultdict(list)
    tk_cot: Dict[Tuple[str, str], dict] = {}
    tk_rong: Dict[str, dict] = {}
    for cate, b in bang.items():
        for code in cot_tt(b):
            if not la_cot_filter(code):
                ten = b["ten"].get(code) or ten_cu.get(code, "")
                tk_cot[(cate, code)] = {"cate": cate, "ten_nh": b.get("ten_nh", ""), "code": code, "ten": ten,
                                        "tong": 0, "so_tron": 0, "vi_du": [], "la_kt": la_cot_kich_thuoc(code, ten),
                                        "goi_y": goi_y_don_vi(code, ten), "da_luu": don_vi_luu.get((cate, code), "")}
        for row in b["rows"]:
            sku, model, variant = row["sku"], row["model"], row["variant"]
            nl = nl_sku.get(sku)
            model_eff = model or (nl.model_code if nl else "")
            cate_pim = row.get("cate_pim") or (nl.category_code if nl else "")
            khoa = sku
            st["so_dong"] += 1
            if not model_eff:
                st["thieu_model"] += 1
                canh_bao.append(["CAO", "Thiếu model_code", sku, cate, "", "Chưa có Mã model"])
            if not cate_pim:
                st["thieu_cate"] += 1
                canh_bao.append(["CAO", "Thiếu category_code", sku, cate, "", "IMPORT chưa có Mã danh mục PIM"])
            e = None
            if co_spec:
                e = cu_sku.get(sku) or cu_mv.get((model_eff, variant))
                if e is None:
                    st["khong_co_pim"] += 1
                    canh_bao.append(["TRUNG BÌNH", "Không có trong file PIM", sku, cate, "", "Không đối chiếu được spec"])
                else:
                    if e["model"] and model_eff and e["model"] != model_eff:
                        st["lech_model"] += 1
                        canh_bao.append(["CAO", "model_code khác file PIM", sku, cate, "",
                                         f"Mới: {model_eff} — PIM: {e['model']}"])
                    if e["cate"] and cate_pim and e["cate"] != cate_pim:
                        st["lech_cate"] += 1
                        canh_bao.append(["CAO", "category_code khác file PIM", sku, cate, "",
                                         f"Mới: {cate_pim} — PIM: {e['cate']}"])
            cu = e["spec"] if e else {}
            dong_ct = []
            for code in cot_tt(b):
                goc = row["vals"].get(code, "")
                if goc and not la_cot_filter(code):
                    kr = khoa_gia_tri_rong(goc)
                    if kr in khoa_rong_qt:
                        x = tk_rong.setdefault(kr, {"khoa": kr, "so_o": 0, "dang": Counter(), "cot": set()})
                        x["so_o"] += 1
                        x["dang"][goc] += 1
                        x["cot"].add(b["ten"].get(code) or code)
                moi, loai = bien_doi_o(cate, khoa, code, goc, sua, don_vi, rong)
                da_sua = loai == "sua"
                if loai:
                    st["bd_" + loai] += 1
                old = cu.get(code, "")
                ten = b["ten"].get(code) or ten_cu.get(code, "")
                t = tk_cot.get((cate, code))
                if t is not None and moi:
                    t["tong"] += 1
                    if la_so_tron(moi):
                        t["so_tron"] += 1
                        if len(t["vi_du"]) < 3 and moi not in t["vi_du"]:
                            t["vi_du"].append(moi)
                    elif len(t["vi_du"]) < 5 and moi[:25] not in t["vi_du"]:
                        t["vi_du"].append(moi[:25])
                if moi and la_cot_filter(code) and _filter_co_chu(moi):
                    st["filter_chu"] += 1
                    canh_bao.append(["CAO", "FILTER có giá trị không phải mã option", sku, cate, code,
                                     f'"{moi}" — PIM chỉ nhận mã số'])
                if moi and t is not None and t["la_kt"] and la_so_tron(moi):
                    st["chua_don_vi"] += 1
                    so_tron_cot[(cate, code)].append(sku)
                if loai == HD_TRONG:
                    tt = TRANG_THAI_BO_QUA
                elif e is None:
                    tt = "" if not moi else "—"
                elif tap(moi, code) == tap(old, code):
                    tt = TRANG_THAI_GIONG if (moi or old) else ""
                elif moi and old:
                    tt = TRANG_THAI_DON_VI if chi_khac_don_vi(moi, old) else TRANG_THAI_KHAC
                elif old:
                    tt = TRANG_THAI_TOOL_TRONG
                else:
                    tt = TRANG_THAI_PIM_TRONG
                if tt in (TRANG_THAI_KHAC, TRANG_THAI_TOOL_TRONG, TRANG_THAI_PIM_TRONG, TRANG_THAI_DON_VI,
                          TRANG_THAI_BO_QUA):
                    st[tt] += 1
                if tt in (TRANG_THAI_KHAC, TRANG_THAI_TOOL_TRONG, TRANG_THAI_PIM_TRONG, TRANG_THAI_DON_VI,
                          TRANG_THAI_BO_QUA) or da_sua:
                    khac.append({"sku": sku, "model": model_eff, "cate": cate, "ma": code, "ten": ten,
                                 "pim_cu": old, "tool_moi": moi, "goc": goc, "trang_thai": tt or TRANG_THAI_GIONG,
                                 "da_sua": da_sua})
                if moi or old or da_sua or goc:
                    dong_ct.append({"ma": code, "ten": ten, "pim_cu": old, "tool_moi": moi, "trang_thai": tt,
                                    "goc": goc, "da_sua": da_sua})
            chi_tiet[khoa] = {"sku": sku, "model": model_eff, "variant": variant, "cate": cate,
                              "tab": b["title"], "co_pim": e is not None, "dong": dong_ct}
    for (cate, code), ds in so_tron_cot.items():
        canh_bao.append(["TRUNG BÌNH", "Kích thước/khối lượng chưa có đơn vị", f"{len(ds)} dòng", cate, code,
                         "VD: " + ", ".join(ds[:8]) + (" ..." if len(ds) > 8 else "")])
    thu_tu = {"CAO": 0, "TRUNG BÌNH": 1, "THÔNG TIN": 2}
    canh_bao.sort(key=lambda x: thu_tu.get(x[0], 9))
    st["rong_tong"] = sum(x["so_o"] for x in tk_rong.values())
    return {
        "co_doi_chieu": co_spec, "stat": dict(st), "khac": khac, "chi_tiet": chi_tiet,
        "canh_bao": pd.DataFrame(canh_bao, columns=["Mức", "Loại", "SKU", "Ngành hàng", "Mã TSKT", "Chi tiết"]),
        "don_vi_cot": [dict(v, vi_du=", ".join(v["vi_du"])) for v in tk_cot.values()],
        "gia_tri_rong": sorted([{"khoa": x["khoa"], "so_o": x["so_o"], "hien": x["dang"].most_common(1)[0][0],
                                 "cac_dang": sorted(x["dang"]), "cot": sorted(x["cot"])} for x in tk_rong.values()],
                               key=lambda d: -d["so_o"]),
    }


# ============================================================================
# KIỂM CHỨNG SKU <-> DATA SP (chỉ ĐỌC, không đổi kết quả map)
# Đi ngược chiều với đối soát: từ ô KẾT QUẢ quay lại DATA SP của CHÍNH SKU đó để chứng minh mỗi giá trị
# lấy đúng từ đúng SKU. Dùng đúng quy tắc của chay_map (mapping theo PROPERTYID, setdefault, quy đổi FILTER).
# ============================================================================
KC_LECH_SKU = "LỆCH SKU"
KC_KHONG_NGUON = "Không thấy trong DATA SP của SKU"
KC_THUOC_TINH_KHAC = "Lấy từ thuộc tính khác (map theo tên)"
KC_SUA_QUY_TAC = "Đã sửa theo quy tắc CMS"
KC_FILTER_THUA = "FILTER có mã không suy ra được từ CMS"
KC_LECH_IMPORT = "model/biến thể khác IMPORT"
KC_LECH_NGANH = "Ngành khác CATEGORYID trong DATA SP"
KC_KY_TU_AN = "Có ký tự ẩn (sẽ tự bỏ khi xuất)"
KC_MUC = {KC_LECH_SKU: "CAO", KC_LECH_IMPORT: "CAO", KC_LECH_NGANH: "CAO", KC_FILTER_THUA: "CAO",
          KC_KHONG_NGUON: "TB", KC_KY_TU_AN: "TB", KC_THUOC_TINH_KHAC: "THAP", KC_SUA_QUY_TAC: "TT"}


def kiem_chung_sku(bang: Dict[str, dict], data_sp: pd.DataFrame, imp: pd.DataFrame, map_tskt: pd.DataFrame,
                   map_filter: pd.DataFrame, opt: dict, quy_doi: Optional[Dict[str, str]] = None,
                   goc_cms: Optional[Dict[str, str]] = None, gia_tri_xuat=None) -> dict:
    """-> {"loi": DataFrame(Mức, Loại, NH, SKU, Mã cột, Giá trị, Nguồn/ghi chú), "tk": {...}}.
    gia_tri_xuat(cate, sku, ma, raw) -> giá trị SẼ XUẤT (để soát ký tự ẩn); None = soát giá trị map."""
    goc_cms = goc_cms or {}
    quy_doi = quy_doi or {}
    cot = ["Mức", "Loại", "NH", "SKU", "Mã cột", "Giá trị", "Nguồn / ghi chú"]
    tk = Counter()
    out: List[list] = []
    if not bang:
        return {"loi": pd.DataFrame(columns=cot), "tk": dict(tk)}
    sku_bang = {r["sku"] for b in bang.values() for r in b["rows"]}
    sp = data_sp[data_sp.PRODUCTCODE.isin(sku_bang)] if data_sp is not None and len(data_sp) else None
    # mapping y hệt chay_map (setdefault: dòng đầu thắng)
    tm: Dict[str, Dict[str, str]] = defaultdict(dict)
    for cate, pid, ma in map_tskt[["cate", "prop_id", "ma"]].itertuples(index=False):
        tm[cate].setdefault(pid, ma)
    fm: Dict[str, Dict[str, str]] = defaultdict(dict)
    for cate, pid, ma in map_filter[["cate", "prop_id", "ma"]].itertuples(index=False):
        fm[cate].setdefault(pid, ma)
    gia_tri_sku: Dict[str, List[Tuple[str, str]]] = defaultdict(list)   # sku -> [(pid, val)]
    cate_sp: Dict[str, str] = {}
    sku_cua_gt: Dict[Tuple[str, str], set] = defaultdict(set)          # (cate, val) -> {sku}
    if sp is not None:
        for sku, pid, val, cid in sp[["PRODUCTCODE", "PROPERTYID", "PROPVALUE", "CATEGORYID"]].itertuples(index=False):
            if cid and sku not in cate_sp:
                cate_sp[sku] = cid
            if pid and val:
                gia_tri_sku[sku].append((pid, val))
    nl = {r.sku: r for r in imp.drop_duplicates("sku").itertuples(index=False)} if imp is not None and len(imp) else {}
    for cate, b in bang.items():
        tmc, fmc = tm.get(cate, {}), fm.get(cate, {})
        for r in b["rows"]:
            for _pid, val in gia_tri_sku.get(r["sku"], ()):
                sku_cua_gt[(cate, val)].add(r["sku"])
        for r in b["rows"]:
            sku = r["sku"]
            tk["sku"] += 1
            n0 = nl.get(sku)
            if n0 is not None and (chuan_hoa_key(n0.model_code) != chuan_hoa_key(r["model"])
                                   or chuan_hoa_key(n0.variant_code) != chuan_hoa_key(r["variant"])):
                out.append([KC_MUC[KC_LECH_IMPORT], KC_LECH_IMPORT, cate, sku, "model_code / variant_code",
                            f"{r['model']} / {r['variant']}", f"IMPORT: {n0.model_code} / {n0.variant_code}"])
            cid = cate_sp.get(sku)
            if cid and cid != cate:
                out.append([KC_MUC[KC_LECH_NGANH], KC_LECH_NGANH, cate, sku, "", "", f"DATA SP CATEGORYID = {cid}"])
            nguon = gia_tri_sku.get(sku, [])
            tat_ca_gt = {v for _, v in nguon}
            for ma in cot_tt(b):
                v = r["vals"].get(ma, "")
                if not v:
                    continue
                tk["o"] += 1
                if gia_tri_xuat is not None:
                    vx = gia_tri_xuat(cate, sku, ma, v)
                    if co_ky_tu_an(vx):
                        tk[KC_KY_TU_AN] += 1
                        out.append([KC_MUC[KC_KY_TU_AN], KC_KY_TU_AN, cate, sku, ma, repr(vx)[:120], "Ký tự ẩn bị loại khi ghi file"])
                g = goc_cms.get(f"{sku}\t{ma}")
                if g is not None:
                    tk[KC_SUA_QUY_TAC] += 1
                    out.append([KC_MUC[KC_SUA_QUY_TAC], KC_SUA_QUY_TAC, cate, sku, ma, v, f"CMS gốc: {g}"])
                    continue
                if la_cot_filter(ma):
                    du_kien = set()
                    for pid, val in nguon:
                        if fmc.get(pid) == ma:
                            oc = quy_doi.get(f"{ma}\t{khoa_quy_doi(val)}") or opt["option_map"].get((ma, khoa_option(val)))
                            if oc:
                                du_kien.add(oc)
                    thua = [x for x in (t.strip() for t in v.split(",")) if x and x not in du_kien]
                    if thua:
                        tk[KC_FILTER_THUA] += 1
                        out.append([KC_MUC[KC_FILTER_THUA], KC_FILTER_THUA, cate, sku, ma, v,
                                    f"Mã không có nguồn: {', '.join(thua)} · suy ra từ CMS: {', '.join(sorted(du_kien)) or '(không)'}"])
                    else:
                        tk["khop"] += 1
                    continue
                dung_cot = {val for pid, val in nguon if tmc.get(pid) == ma}
                for phan in [x for x in v.split(SEP_TSKT) if x]:
                    if phan in dung_cot:
                        tk["khop"] += 1
                    elif phan in tat_ca_gt:
                        tk[KC_THUOC_TINH_KHAC] += 1
                        out.append([KC_MUC[KC_THUOC_TINH_KHAC], KC_THUOC_TINH_KHAC, cate, sku, ma, phan,
                                    "Giá trị có trong DATA SP của SKU nhưng ở thuộc tính khác"])
                    else:
                        khac = sorted(sku_cua_gt.get((cate, phan), set()) - {sku})
                        if khac:
                            tk[KC_LECH_SKU] += 1
                            out.append([KC_MUC[KC_LECH_SKU], KC_LECH_SKU, cate, sku, ma, phan,
                                        "Giá trị chỉ có ở SKU khác: " + ", ".join(khac[:5])])
                        else:
                            tk[KC_KHONG_NGUON] += 1
                            out.append([KC_MUC[KC_KHONG_NGUON], KC_KHONG_NGUON, cate, sku, ma, phan,
                                        "Không có trong DATA SP của SKU này"])
    thu_tu = {"CAO": 0, "TB": 1, "THAP": 2, "TT": 3}
    out.sort(key=lambda x: (thu_tu.get(x[0], 9), x[1], x[2], x[3]))
    return {"loi": pd.DataFrame(out, columns=cot), "tk": dict(tk)}


# ============================================================================
# QC NGẦM: NHẤT QUÁN NGÀNH HÀNG (cấu hình · DATA SP · mapping TSKT · mapping FILTER · DATA PIM)
# Chỉ ĐỌC. Không đổi dữ liệu, không đổi kết quả map. Dùng đúng quy tắc của chay_map (setdefault, la_cot_filter…).
# ============================================================================
def kiem_tra_nhat_quan(cau_hinh: Dict[str, dict], data_sp: Optional[pd.DataFrame], map_tskt: pd.DataFrame,
                       map_filter: pd.DataFrame, opt: dict, quy_doi: Optional[Dict[str, str]] = None,
                       chi_cate: Optional[List[str]] = None) -> dict:
    """-> {"loi": DataFrame(Mức, NH, Vùng, Vấn đề, Số, Chi tiết), "nganh": DataFrame tóm tắt từng ngành}.
    chi_cate=None: kiểm các ngành có trong DATA SP (+ ngành có cấu hình nếu DATA SP trống)."""
    quy_doi = quy_doi or {}
    sp = data_sp if data_sp is not None and len(data_sp) else pd.DataFrame(columns=COT_DATA_SP)
    cate_sp = sorted(set(sp.CATEGORYID) - {""}) if len(sp) else []
    if chi_cate is None:
        chi_cate = cate_sp or sorted(cau_hinh)
    chi_cate = [c for c in dict.fromkeys(chi_cate) if c]
    tm_rows = map_tskt[map_tskt.cate.isin(chi_cate)] if len(map_tskt) else map_tskt
    fm_rows = map_filter[map_filter.cate.isin(chi_cate)] if len(map_filter) else map_filter
    opt_ds = opt.get("opt_ds", {}) if opt else {}
    loi: List[list] = []
    tom: List[dict] = []

    def them(muc, cate, vung, van_de, so, chi_tiet=""):
        loi.append([muc, cate, vung, van_de, int(so), chi_tiet])
    for cate in chi_cate:
        cfg = cau_hinh.get(cate) or {}
        cot = [c for c in (cfg.get("cot") or []) if c.lower() not in COT_KHONG_PHAI_SPEC]
        cot_set = set(cot)
        ten_nh = cfg.get("ten", "")
        t_c = tm_rows[tm_rows.cate == cate] if len(tm_rows) else tm_rows
        f_c = fm_rows[fm_rows.cate == cate] if len(fm_rows) else fm_rows
        if not ten_nh:
            ten_nh = next((x for x in list(t_c.cate_name) + list(f_c.cate_name) if x), "")
        sp_c = sp[sp.CATEGORYID == cate] if len(sp) else sp
        n_sku = sp_c.PRODUCTCODE.nunique() if len(sp_c) else 0
        # mapping theo đúng chay_map: setdefault (dòng đầu thắng)
        tm: Dict[str, str] = {}
        tm_tat: Dict[str, set] = defaultdict(set)
        for pid, ma in t_c[["prop_id", "ma"]].itertuples(index=False):
            tm.setdefault(pid, ma)
            tm_tat[pid].add(ma)
        fm: Dict[str, str] = {}
        fm_tat: Dict[str, set] = defaultdict(set)
        for pid, ma in f_c[["prop_id", "ma"]].itertuples(index=False):
            fm.setdefault(pid, ma)
            fm_tat[pid].add(ma)
        ma_map = set(tm.values()) | set(fm.values())
        # A/B: sót cấu hình / sót mapping
        if n_sku and not cot:
            them("CAO", cate, "Cấu hình", "Ngành có trong DATA SP nhưng CHƯA có cấu hình cột", n_sku,
                 f"{n_sku} SKU sẽ bị bỏ qua khi map — thêm cấu hình ngành (file mẫu ngành / ⚙️ Cấu hình)")
        if n_sku and not tm and not fm:
            them("CAO", cate, "Mapping", "Ngành có trong DATA SP nhưng CHƯA có mapping TSKT/FILTER", n_sku,
                 f"{n_sku} SKU sẽ bị bỏ qua khi map (giống desktop) — nạp MAPPING TSKT/FILTER của ngành")
        if cot and not tm and not fm and not n_sku:
            them("TB", cate, "Mapping", "Ngành có cấu hình nhưng chưa có mapping nào", len(cot),
                 "Khi có DATA SP của ngành này sẽ không map được")
        # C: cột cấu hình không có mapping nào trỏ tới
        if cot and (tm or fm):
            trong = [c for c in cot if c not in ma_map]
            if trong:
                them("TB", cate, "Cấu hình ↔ Mapping", "Cột cấu hình chưa có mapping nào trỏ tới (luôn trống)",
                     len(trong), ", ".join(trong[:15]) + (" …" if len(trong) > 15 else ""))
        # D: mapping trỏ tới cột không có trong cấu hình
        if cot:
            ngoai = sorted(m for m in ma_map if m and m not in cot_set)
            if ngoai:
                them("CAO", cate, "Mapping ↔ Cấu hình", "Mapping trỏ tới cột KHÔNG có trong cấu hình (giá trị bị bỏ)",
                     len(ngoai), ", ".join(ngoai[:15]) + (" …" if len(ngoai) > 15 else ""))
        # G: 1 PROPERTYID map tới nhiều cột
        for vung, tat, dung in (("Mapping TSKT", tm_tat, tm), ("Mapping FILTER", fm_tat, fm)):
            trung = {pid: sorted(v) for pid, v in tat.items() if len(v) > 1}
            if trung:
                them("TB", cate, vung, "1 PROPERTYID trỏ tới nhiều cột — chỉ dòng đầu được dùng", len(trung),
                     "; ".join(f"{pid}: dùng {dung[pid]} (bỏ {', '.join(x for x in v if x != dung[pid])})"
                               for pid, v in list(trung.items())[:8]))
        # I: sai loại cột
        sai_f = sorted({m for m in fm.values() if m and not la_cot_filter(m)})
        if sai_f:
            them("CAO", cate, "Mapping FILTER", "Mapping FILTER trỏ tới cột KHÔNG phải FILTER", len(sai_f),
                 ", ".join(sai_f[:10]))
        sai_t = sorted({m for m in tm.values() if m and la_cot_filter(m)})
        if sai_t:
            them("CAO", cate, "Mapping TSKT", "Mapping TSKT trỏ tới cột FILTER (cột cần MÃ option)", len(sai_t),
                 ", ".join(sai_t[:10]))
        # J: cột FILTER không có option trong DATA PIM
        cot_f = sorted({m for m in list(fm.values()) + cot if m and la_cot_filter(m)})
        thieu_opt = [m for m in cot_f if not opt_ds.get(m)]
        if thieu_opt:
            them("CAO", cate, "DATA PIM", "Cột FILTER không có option nào trong DATA PIM (mọi ô sẽ trống)",
                 len(thieu_opt), ", ".join(thieu_opt[:10]))
        # L: cột cấu hình lặp
        lap = sorted({c for c in cot if cot.count(c) > 1})
        if lap:
            them("TB", cate, "Cấu hình", "Cột bị lặp trong cấu hình", len(lap), ", ".join(lap))
        # E/K: theo DATA SP của ngành
        n_pid = n_pid_map = 0
        if len(sp_c):
            pid_ten: Dict[str, str] = {}
            pid_dong: Counter = Counter()
            for pid, pname in sp_c[["PROPERTYID", "PROPERTYNAME"]].itertuples(index=False):
                if pid:
                    pid_dong[pid] += 1
                    pid_ten.setdefault(pid, pname)
            n_pid = len(pid_dong)
            chua = [pid for pid in pid_dong if pid not in tm and pid not in fm]
            n_pid_map = n_pid - len(chua)
            if chua and (tm or fm):
                # gợi ý (chỉ gợi ý): tên thuộc tính CMS trùng tên 1 cột cấu hình CHƯA có mapping
                ten_cot = cfg.get("ten_cot") or {}
                cot_trong = {chuan_hoa_ten(ten_cot.get(c, "")): c for c in cot if c not in ma_map and ten_cot.get(c)}
                def _gy(pid):
                    m = cot_trong.get(chuan_hoa_ten(pid_ten.get(pid, "")))
                    return f" → gợi ý cột {m}" if m else ""
                them("TB", cate, "DATA SP ↔ Mapping", "Thuộc tính CMS trong DATA SP chưa có mapping (không vào file)",
                     len(chua), "; ".join(f"{pid} {pid_ten.get(pid, '')} ({pid_dong[pid]} dòng){_gy(pid)}"
                                          for pid in sorted(chua, key=lambda x: -pid_dong[x])[:10]))
            # giá trị FILTER không khớp option (giống chay_map: quy đổi trước, rồi option_map)
            khong_khop: Dict[Tuple[str, str], int] = Counter()
            om = opt.get("option_map", {}) if opt else {}
            for pid, val in sp_c[["PROPERTYID", "PROPVALUE"]].itertuples(index=False):
                ma = fm.get(pid)
                if not ma or not val or ma not in cot_set:
                    continue
                if not (quy_doi.get(f"{ma}\t{khoa_quy_doi(val)}") or om.get((ma, khoa_option(val)))):
                    khong_khop[(ma, val)] += 1
            if khong_khop:
                them("TB", cate, "DATA SP ↔ DATA PIM", "Giá trị FILTER trong DATA SP không khớp option DATA PIM",
                     len(khong_khop), "; ".join(f"{m}=\"{v}\"" for (m, v) in list(khong_khop)[:8]))
            # F: mapping có PROPERTYID không xuất hiện trong lô (thông tin)
            khong_co = [pid for pid in list(tm) + list(fm) if pid not in pid_dong]
            if khong_co:
                them("TT", cate, "Mapping ↔ DATA SP", "PROPERTYID có mapping nhưng lô này không có dữ liệu",
                     len(set(khong_co)), "Bình thường nếu lô ít thông số")
        so_cao = sum(1 for x in loi if x[1] == cate and x[0] == "CAO")
        so_tb = sum(1 for x in loi if x[1] == cate and x[0] == "TB")
        tom.append({"Mã NH": cate, "Tên ngành": ten_nh, "SKU (DATA SP)": n_sku, "Dòng DATA SP": len(sp_c),
                    "Cột cấu hình": len(cot), "Cột có mapping": len([c for c in cot if c in ma_map]),
                    "% cột có mapping": round(100 * len([c for c in cot if c in ma_map]) / len(cot), 1) if cot else 0.0,
                    "Thuộc tính CMS đã map": f"{n_pid_map}/{n_pid}" if n_pid else "",
                    "Mapping TSKT": len(t_c), "Mapping FILTER": len(f_c),
                    "Trạng thái": "⛔ Lỗi" if so_cao else ("⚠️ Cần xem" if so_tb else "✅ Ổn")})
    thu_tu = {"CAO": 0, "TB": 1, "THAP": 2, "TT": 3}
    loi.sort(key=lambda x: (thu_tu.get(x[0], 9), x[1]))
    return {"loi": pd.DataFrame(loi, columns=["Mức", "NH", "Vùng", "Vấn đề", "Số", "Chi tiết"]),
            "nganh": pd.DataFrame(tom)}


def giai_nghia_filter(code: str, gia_tri: str, opt: dict) -> str:
    if not la_cot_filter(code) or not gia_tri:
        return ""
    out = []
    for t in re.split(r"[,|]", lam_sach_gia_tri_pim(gia_tri, code)):
        t = chuan_hoa_id(t.strip())
        if t:
            out.append(f"{t}={opt['opt_ten'].get((code, t), '❓KHÔNG CÓ trong DATA PIM')}")
    return " · ".join(out)


# ============================================================================
# §7 XUẤT FILE IMPORT
# ============================================================================
def _o_sach(v) -> str:
    """Chốt chặn cuối trước khi ghi ô vào file import: bỏ ký tự ẩn (kể cả trong ô SỬA TAY / dán từ web).
    Ô không có ký tự ẩn giữ NGUYÊN y hệt (không trim) -> file import không đổi so với desktop."""
    t = str(v)
    t2 = bo_ky_tu_an(t)
    return t2.strip() if t2 != t else t


def _xlsx_text(rows: List[list]) -> bytes:
    """File import: sheet "Export Product Template", MỌI ô dạng Text "@" (cả cột), 2 dòng tiêu đề in đậm, cố định A3.
    Dùng xlsxwriter (nhanh ~4 lần, ít RAM) nếu có; không thì openpyxl — nội dung ô y hệt nhau."""
    try:
        import xlsxwriter
    except ImportError:
        xlsxwriter = None
    n = max((len(r) for r in rows), default=0)
    if xlsxwriter is not None:
        buf = io.BytesIO()
        wb = xlsxwriter.Workbook(buf, {"in_memory": True, "strings_to_numbers": False, "strings_to_formulas": False,
                                       "strings_to_urls": False})
        ws = wb.add_worksheet("Export Product Template")
        f_txt = wb.add_format({"num_format": "@"})
        f_dam = wb.add_format({"num_format": "@", "bold": True})
        if n:
            ws.set_column(0, n - 1, None, f_txt)
        for i, r in enumerate(rows):
            f = f_dam if i <= 1 else f_txt
            for j, v in enumerate(r):
                if v not in (None, ""):
                    v = _o_sach(v)
                    if v:
                        ws.write_string(i, j, v, f)
        ws.freeze_panes(2, 0)
        wb.close()
        return buf.getvalue()
    wb = Workbook()
    ws = wb.active
    ws.title = "Export Product Template"
    for c in range(1, n + 1):
        ws.column_dimensions[get_column_letter(c)].number_format = "@"
    bold = Font(bold=True)
    for i, r in enumerate(rows, start=1):
        for j, v in enumerate(r, start=1):
            if v not in (None, "") and _o_sach(v):
                cell = ws.cell(row=i, column=j, value=_o_sach(v))
                cell.number_format = "@"
                if i <= 2:
                    cell.font = bold
    ws.freeze_panes = "A3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xuat_file_import(bang: Dict[str, dict], imp: pd.DataFrame, sua: dict, don_vi: dict, rong: dict,
                     bo_cot_sku: bool = True, chi_cate: Optional[List[str]] = None,
                     bo_dong_trong: bool = False) -> dict:
    """-> {"files": [(tên, bytes, số dòng)], "zip": bytes, "so_o_sua", "so_o_dv", "so_o_rong", "bo_dong", "bo_trong"}.
    bo_dong_trong=True: SKU không có giá trị thuộc tính nào -> KHÔNG đưa vào file import (đưa vào file xin data).
    Mặc định False = như desktop (vẫn xuất dòng trống)."""
    stamp = datetime.now(VN_TZ).strftime("%Y%m%d_%H%M")
    sku_imp = set(imp.sku) if imp is not None and len(imp) else set()
    files: List[Tuple[str, bytes, int]] = []
    dem = Counter()
    for cate, b in bang.items():
        if chi_cate and cate not in chi_cate:
            continue
        cot = ["model_code"] + ([] if bo_cot_sku else ["sku"]) + ["variant_code"] + b["attr"]
        ten = {"model_code": "Mã model", "sku": "Mã sản phẩm ERP", "variant_code": "Mã biến thể"}
        h2 = [ten.get(c) or b["ten"].get(c, "") for c in cot]
        model_rows, var_rows = [], []
        for r in b["rows"]:
            if sku_imp and r["sku"] not in sku_imp:
                dem["bo_dong"] += 1
                continue
            out = []
            for c in cot:
                if c in ("model_code", "sku", "variant_code"):
                    out.append(r["vals"].get(c, ""))
                    continue
                v, loai = bien_doi_o(cate, r["sku"], c, r["vals"].get(c, ""), sua, don_vi, rong)
                if loai:
                    dem[{"sua": "so_o_sua", "dv": "so_o_dv", "bd": "so_o_bd"}.get(loai, "so_o_rong")] += 1
                out.append(v)
            if bo_dong_trong and not any(v for c, v in zip(cot, out) if c not in ("model_code", "sku", "variant_code")):
                dem["bo_trong"] += 1
                continue
            (var_rows if chuan_hoa_key(r["vals"].get("variant_code")) else model_rows).append(out)
        nhan = ten_tab_desktop(b["title"])
        for hau, data in (("_MODEL", model_rows), ("_BIENTHE", var_rows)):
            if data:
                files.append((f"PIM_{nhan}{hau}_{stamp}.xlsx", _xlsx_text([cot, h2] + data), len(data)))
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w", zipfile.ZIP_DEFLATED) as zf:
        for t, d, _ in files:
            zf.writestr(t, d)
        if sua:
            zf.writestr("_CHINH_SUA_TAY.txt", "\r\n".join(
                ["NGÀNH HÀNG\tSKU\tMÃ TSKT\tGIÁ TRỊ SỬA TAY"] +
                [f"{c}\t{k}\t{m}\t{v}" for (c, k, m), v in sorted(sua.items())]).encode("utf-8-sig"))
    return {"files": files, "zip": zbuf.getvalue(), "stamp": stamp, **dem}


def xlsx_nhieu_sheet(sheets: Dict[str, pd.DataFrame]) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for ten, df in sheets.items():
            df.to_excel(w, sheet_name=ten[:31], index=False)
    return buf.getvalue()


# ============================================================================
# §8 KIỂM TRA THÔNG MINH (không cần AI — miễn phí, chạy tức thì)
# ============================================================================
# Dò các lỗi "trông thì hợp lệ nhưng sai" mà so spec PIM không bắt được:
#   - giá trị bất thường so với cả cột (99 kg trong cột toàn ~1 kg)
#   - cột trộn đơn vị (472 mm giữa cột toàn cm) -> gợi ý quy đổi
#   - cùng 1 chữ viết nhiều kiểu (không / Không / KHÔNG / Không.) -> gợi ý kiểu phổ biến
#   - lỗi gõ: "18. 4 cm", "56.3 Kg", dấu chấm/phẩy thừa cuối, "A|A" trùng, "|" rỗng
# Mọi gợi ý chỉ là ĐỀ XUẤT — người dùng tick mới thành sửa tay.
_SO_DV_TACH = re.compile(r"^\s*(\d+(?:[.,]\d+)?)\s*(mm|cm|m|kg|g|inch)?\s*$", re.IGNORECASE)
_QUY_DOI = {("mm", "cm"): 0.1, ("m", "cm"): 100.0, ("cm", "mm"): 10.0, ("g", "kg"): 0.001, ("kg", "g"): 1000.0,
            ("cm", "m"): 0.01}
MUC_CAO, MUC_TB, MUC_THAP = "Cao", "Trung bình", "Thấp"


def _so(s: str) -> Optional[Tuple[float, str]]:
    m = _SO_DV_TACH.match(s or "")
    if not m:
        return None
    return float(m.group(1).replace(",", ".")), (m.group(2) or "").lower()


def _fmt_so(x: float) -> str:
    s = f"{x:.3f}".rstrip("0").rstrip(".")
    return s or "0"


def _chuan_kieu_viet(s: str) -> str:
    return unicodedata.normalize("NFC", _WS_RE.sub(" ", s)).strip(" .,;").lower()


_DOAN_SO_DV = re.compile(r"^(\d+)(?:\.\s+(\d+)|([.,]\d+))?\s*(kg|cm|mm|g|m|inch)$", re.IGNORECASE)


def sua_loi_go(v: str) -> Tuple[str, List[str]]:
    """Sửa lỗi gõ CHẮC CHẮN (không đổi nghĩa) -> (giá trị mới, các lý do).
    Đơn vị chỉ chuẩn hoá khi CẢ ĐOẠN là số + đơn vị (không đụng chữ tự do như 'Driver 40mm')."""
    ly_do: List[str] = []
    doan_ra = []
    parts = v.split("|") if "|" in v else [v]
    for p in parts:
        p0 = p
        p = re.sub(r"\s{2,}", " ", p).strip()
        m = _DOAN_SO_DV.match(p)
        if m:
            so = m.group(1) + ("." + m.group(2) if m.group(2) else (m.group(3) or ""))
            p_moi = f"{so} {m.group(4).lower()}"
            if m.group(2):
                ly_do.append("dấu thập phân bị tách (18. 4)")
            if m.group(4) != m.group(4).lower():
                ly_do.append("đơn vị viết hoa")
            if p_moi != p and not m.group(2) and m.group(4) == m.group(4).lower():
                ly_do.append("thiếu khoảng trắng giữa số và đơn vị")
            p = p_moi
        elif p != p0.strip():
            ly_do.append("khoảng trắng thừa")
        p2 = re.sub(r"[ ,;]+$", "", p)
        if p2.endswith(".") and not re.search(r"\d\.$", p2) and len(p2) <= 40:
            p2 = p2.rstrip(". ")
        if p2 != p:
            ly_do.append("dấu câu thừa cuối")
        if p2:
            doan_ra.append(p2)
        elif p0.strip() == "" and "|" in v:
            ly_do.append("đoạn rỗng giữa dấu |")
    uniq = list(dict.fromkeys(doan_ra))
    if len(uniq) < len(doan_ra):
        ly_do.append("giá trị lặp lại trong danh sách |")
    return "|".join(uniq), list(dict.fromkeys(ly_do))


def kiem_tra_thong_minh(bang: Dict[str, dict], sua: dict, don_vi: dict, rong: dict) -> pd.DataFrame:
    """-> DataFrame gợi ý: cate, sku, ma, ten, gia_tri, goi_y, loai, ly_do, muc_do."""
    out: List[dict] = []

    def them(cate, sku, ma, ten, v, goi_y, loai, ly_do, muc):
        out.append({"cate": cate, "sku": sku, "ma": ma, "ten": ten, "gia_tri": v, "goi_y": goi_y,
                    "loai": loai, "ly_do": ly_do, "muc_do": muc})
    for cate, b in bang.items():
        for ma in cot_tt(b):
            if la_cot_filter(ma):
                continue
            ten = b["ten"].get(ma, "")
            gt: List[Tuple[str, str]] = []
            for r in b["rows"]:
                v, _ = bien_doi_o(cate, r["sku"], ma, r["vals"].get(ma, ""), sua, don_vi, rong)
                if v:
                    gt.append((r["sku"], v))
            if not gt:
                continue
            da_goi_y: set = set()
            # --- 1. lỗi gõ chắc chắn
            for sku, v in gt:
                s2, ld = sua_loi_go(v)
                if s2 != v and ld:
                    chi_khoang_trang = set(ld) <= {"khoảng trắng thừa", "thiếu khoảng trắng giữa số và đơn vị"}
                    them(cate, sku, ma, ten, v, s2, "Lỗi gõ / định dạng", "; ".join(ld).capitalize(),
                         MUC_THAP if chi_khoang_trang else MUC_TB)
                    da_goi_y.add(sku)
            # --- 2. cùng chữ nhiều kiểu viết (chỉ với giá trị ngắn, không số)
            nhom: Dict[str, Counter] = defaultdict(Counter)
            for _, v in gt:
                if len(v) <= 40 and not re.search(r"\d", v):
                    nhom[_chuan_kieu_viet(v)][v] += 1
            chuan = {k: c.most_common(1)[0][0] for k, c in nhom.items() if len(c) > 1}
            for sku, v in gt:
                k = _chuan_kieu_viet(v)
                if k in chuan and v != chuan[k] and sku not in da_goi_y:
                    an = unicodedata.normalize("NFC", v) == unicodedata.normalize("NFC", chuan[k])
                    them(cate, sku, ma, ten, v, chuan[k], "Nhiều kiểu viết",
                         ("Nhìn giống hệt nhưng khác mã chữ tiếng Việt (gõ tổ hợp/dựng sẵn) — dễ làm lọc/so khớp sai"
                          if an else f'{sum(nhom[k].values())} ô cùng nghĩa trong cột, kiểu phổ biến nhất là "{chuan[k]}"'),
                         MUC_TB if an else MUC_THAP)
            # --- 3. đơn vị lẫn lộn + giá trị bất thường (cột số)
            so = [(sku, v, _so(v)) for sku, v in gt]
            so = [(sku, v, x) for sku, v, x in so if x]
            if len(so) < 8 or len(so) < 0.5 * len(gt):
                continue
            dv_dem = Counter(x[1] for _, _, x in so if x[1])
            dv_chinh = dv_dem.most_common(1)[0][0] if dv_dem else ""
            for sku, v, (n, dv) in so:
                if dv and dv_chinh and dv != dv_chinh and dv_dem[dv] <= 0.2 * len(so) and (dv, dv_chinh) in _QUY_DOI:
                    goi_y = f"{_fmt_so(n * _QUY_DOI[(dv, dv_chinh)])} {dv_chinh}"
                    them(cate, sku, ma, ten, v, goi_y, "Lẫn đơn vị",
                         f"Cột chủ yếu dùng {dv_chinh} ({dv_dem[dv_chinh]} ô) — ô này dùng {dv}", MUC_TB)
                    da_goi_y.add(sku)
            # quy về đơn vị chính để so độ lớn
            gia_tri_chuan = []
            for sku, v, (n, dv) in so:
                if dv and dv_chinh and dv != dv_chinh and (dv, dv_chinh) in _QUY_DOI:
                    n = n * _QUY_DOI[(dv, dv_chinh)]
                gia_tri_chuan.append((sku, v, n))
            ds = sorted(n for _, _, n in gia_tri_chuan if n > 0)
            if len(ds) < 8:
                continue
            trung_vi = ds[len(ds) // 2]
            if trung_vi <= 0:
                continue
            for sku, v, n in gia_tri_chuan:
                if sku in da_goi_y or n <= 0:
                    continue
                ti_le = n / trung_vi
                if ti_le >= 25 or ti_le <= 1 / 25:
                    goi_y = ""
                    x = _so(v)
                    if x and ti_le >= 25:
                        for he_so in (10, 100, 1000):
                            if 0.2 <= ti_le / he_so <= 5:
                                goi_y = f"{_fmt_so(x[0] / he_so)}{(' ' + x[1]) if x[1] else ''}"
                                break
                    dvc = f" {dv_chinh}" if dv_chinh else ""
                    them(cate, sku, ma, ten, v, goi_y, "Giá trị bất thường",
                         (f"Lớn hơn ~{ti_le:.0f}× trung vị cột ({_fmt_so(trung_vi)}{dvc}) — có thể sai dấu thập phân / đơn vị"
                          if ti_le >= 1 else
                          f"Nhỏ hơn ~{1 / ti_le:.0f}× trung vị cột ({_fmt_so(trung_vi)}{dvc}) — kiểm tra lại"),
                         MUC_CAO if ti_le >= 1 else MUC_TB)
    df = pd.DataFrame(out, columns=["cate", "sku", "ma", "ten", "gia_tri", "goi_y", "loai", "ly_do", "muc_do"])
    thu_tu = {MUC_CAO: 0, MUC_TB: 1, MUC_THAP: 2}
    return df.sort_values(["muc_do", "ma"], key=lambda s: s.map(thu_tu) if s.name == "muc_do" else s,
                          ignore_index=True) if len(df) else df



# ============================================================================
# §9 ĐỐI SOÁT CMS -> KẾT QUẢ (báo mất dữ liệu khi map + nhiều trường hợp khác)
# ============================================================================
# Đi từng giá trị CMS (DATA SP) của lô, xác định giá trị đó CÓ vào file import
# không; nếu không -> vì sao + gợi ý cách sửa. Kèm độ phủ từng cột.
import difflib  # noqa: E402

LOAI_DS = {
    "filter": "FILTER: giá trị CMS không có option",
    "cot": "Mapping trỏ tới cột KHÔNG có trong cấu hình",
    "nganh_khac": "Thuộc tính có mapping ở ngành khác, thiếu ở ngành này",
    "chua_map": "Thuộc tính CMS chưa có mapping",
    "mat": "Giá trị CMS không xuất hiện trong kết quả",
    "khong_data": "SKU không có dữ liệu CMS",
    "it_thong_so": "SKU thiếu nhiều thông số so với cùng ngành",
    "cot_trong": "Cột luôn trống cả lô",
    "lech_ft": "Lệch FILTER ↔ TSKT cùng loại",
    "gop": "Nhiều thuộc tính CMS gộp vào 1 cột",
}


@lru_cache(maxsize=_CACHE)
def khoa_quy_doi(v: str) -> str:
    return unicodedata.normalize("NFC", _WS_RE.sub(" ", chuan_hoa_key(v))).strip(" .").lower()


def goi_y_option(ma: str, val: str, opt: dict, n: int = 3) -> List[Tuple[str, str]]:
    """Option gần giống nhất (theo tên) cho giá trị CMS không khớp."""
    ds = opt["opt_ds"].get(ma, [])
    if not ds:
        return []
    ten = {tn.lower(): (oc, tn) for oc, tn in ds}
    k = khoa_quy_doi(val)
    gan = difflib.get_close_matches(k, list(ten), n=n, cutoff=0.55)
    # thêm: option nằm trọn trong giá trị CMS hoặc ngược lại ("Wi-Fi 6" ~ "Wi-Fi")
    for t in ten:
        if len(gan) >= n:
            break
        if t not in gan and len(t) >= 3 and (t in k or (len(k) >= 3 and k in t)):
            gan.append(t)
    return [ten[t] for t in gan[:n]]


def doi_soat_map(data_sp: pd.DataFrame, imp: pd.DataFrame, bang: Dict[str, dict], cau_hinh: Dict[str, dict],
                 map_tskt: pd.DataFrame, map_filter: pd.DataFrame, opt: dict,
                 quy_doi: Optional[Dict[str, str]] = None, spec: Optional[pd.DataFrame] = None,
                 sua_sku: Optional[dict] = None, sua_gt: Optional[dict] = None) -> dict:
    quy_doi = quy_doi or {}
    _gy: Dict[Tuple[str, str], list] = {}
    o_co_quy_tac = {tuple(k.split("\t")) for k in (sua_sku or {})}
    cot_co_quy_tac = {tuple(k.split("\t")[:2]) for k in (sua_gt or {})}
    sku_cate = {r["sku"]: c for c, b in bang.items() for r in b["rows"]}
    row_of = {r["sku"]: r for b in bang.values() for r in b["rows"]}
    tm: Dict[str, Dict[str, str]] = defaultdict(dict)
    for c, pid, ma in map_tskt[["cate", "prop_id", "ma"]].itertuples(index=False):
        tm[c].setdefault(pid, ma)
    fm: Dict[str, Dict[str, str]] = defaultdict(dict)
    for c, pid, ma in map_filter[["cate", "prop_id", "ma"]].itertuples(index=False):
        fm[c].setdefault(pid, ma)
    nganh_khac: Dict[str, List[Tuple[str, str, str]]] = defaultdict(list)  # pid -> [(loai, cate, ma)]
    for loai, d in (("tskt", tm), ("filter", fm)):
        for c, m in d.items():
            for pid, ma in m.items():
                nganh_khac[pid].append((loai, c, ma))
    nhom: Dict[tuple, dict] = {}

    def them(muc, loai, cate, pid="", pname="", val="", ma="", sku="", goi_y="", extra=None):
        k = (loai, cate, pid, ma, val if loai == "filter" else "")
        g = nhom.get(k)
        if g is None:
            g = nhom[k] = {"Mức": muc, "Loại": LOAI_DS[loai], "_loai": loai, "Ngành": cate, "Mã thuộc tính CMS": pid,
                           "Tên thuộc tính CMS": pname, "Giá trị CMS": val, "Mã cột": ma, "Số SKU": 0,
                           "SKU ví dụ": [], "Gợi ý": goi_y, **(extra or {})}
        if sku and sku not in g["SKU ví dụ"]:
            g["Số SKU"] += 1
            if len(g["SKU ví dụ"]) < 5:
                g["SKU ví dụ"].append(sku)
    sp = data_sp[data_sp.PRODUCTCODE.isin(set(sku_cate))] if len(data_sp) else data_sp
    nguon_cot: Dict[Tuple[str, str, str], set] = defaultdict(set)  # (cate, sku, ma) -> prop ids
    cot_cua = {c: set(b["attr"]) for c, b in bang.items()}
    for sku, pid, pname, val in sp[["PRODUCTCODE", "PROPERTYID", "PROPERTYNAME", "PROPVALUE"]].itertuples(index=False):
        if not pid or not val:
            continue
        cate = sku_cate[sku]
        cot = cot_cua[cate]
        vals = row_of[sku]["vals"]
        da = False
        for loai, d in (("tskt", tm), ("filter", fm)):
            ma = d.get(cate, {}).get(pid)
            if not ma:
                continue
            da = True
            if ma not in cot:
                them(MUC_CAO, "cot", cate, pid, pname, "", ma, sku,
                     f"Thêm cột {ma} vào CẤU HÌNH ngành {cate} (hoặc sửa mapping nếu trỏ nhầm)")
                continue
            nguon_cot[(cate, sku, ma)].add(pid)
            if loai == "filter":
                oc = quy_doi.get(f"{ma}\t{khoa_quy_doi(val)}") or opt["option_map"].get((ma, khoa_option(val)))
                if not oc:
                    if khoa_gia_tri_rong(val) in _KHOA_RONG_MAC_DINH:
                        them(MUC_THAP, "filter", cate, pid, pname, val, ma, sku,
                             "Giá trị kiểu Không/Đang cập nhật — để trống FILTER là hợp lý", {"_goi_y": []})
                        continue
                    gy = _gy.get((ma, val))
                    if gy is None:  # gợi ý option chỉ tính 1 lần cho mỗi (cột, giá trị) — không phải mỗi SKU
                        gy = _gy[(ma, val)] = goi_y_option(ma, val, opt)
                    them(MUC_CAO, "filter", cate, pid, pname, val, ma, sku,
                         " · ".join(f"{a}={b}" for a, b in gy) or "Không thấy option gần giống — thêm option trên PIM",
                         {"_goi_y": gy})
            else:
                if (sku, ma) in o_co_quy_tac or (cate, ma) in cot_co_quy_tac:
                    continue  # ô đã được sửa theo đề xuất -> khác CMS là đúng ý
                if _gon(val) not in _gon(str(vals.get(ma, ""))):
                    them(MUC_TB, "mat", cate, pid, pname, val, ma, sku, "Kiểm tra lại mapping / sửa tay")
        if not da:
            khac = [(lo, c, m) for lo, c, m in nganh_khac.get(pid, []) if c != cate and m in cot]
            # đã được map dự phòng theo tên vào 1 cột? -> bỏ qua
            if any(val in str(vals.get(m, "")).split(SEP_TSKT) for m in cot):
                continue
            if khac:
                lo, c2, m2 = khac[0]
                them(MUC_TB, "nganh_khac", cate, pid, pname, "", m2, sku,
                     f"Ngành {c2} map thuộc tính này → {m2} ({lo.upper()}); copy sang ngành {cate}",
                     {"_copy": (lo, m2)})
            else:
                ref = map_theo_ten().get(cate, {}).get(chuan_hoa_ten(pname))
                if ref and ref in cot:
                    them(MUC_TB, "chua_map", cate, pid, pname, val[:60], ref, sku,
                         f"Bảng tham chiếu tên gợi ý map → {ref} (kiểm tra lại rồi thêm ở tab 🧩)")
                else:
                    them(MUC_THAP, "chua_map", cate, pid, pname, val[:60], "", sku, "Thêm vào MAPPING nếu cần (tab 🧩)")
    # --- SKU không có dữ liệu / thiếu nhiều thông số
    co_data = Counter(sp.PRODUCTCODE) if len(sp) else Counter()
    for cate, b in bang.items():
        ctt = cot_tt(b)
        dem = {r["sku"]: sum(1 for m in ctt if r["vals"].get(m)) for r in b["rows"]}
        ds = sorted(dem.values())
        tv = ds[len(ds) // 2] if ds else 0
        for sku, n in dem.items():
            if sku not in co_data:
                them(MUC_CAO, "khong_data", cate, sku=sku, goi_y="Nạp thêm file CMS export có SKU này")
            elif tv >= 6 and n < 0.35 * tv:
                them(MUC_CAO if n == 0 else MUC_TB, "it_thong_so", cate, sku=sku, ma=f"{n}/{tv} cột",
                     goi_y=(f"Có {co_data[sku]} dòng CMS nhưng KHÔNG thuộc tính nào map được — kiểm tra mã thuộc tính "
                            "(có thể khác bộ thuộc tính TGDĐ/ĐMX) hoặc ngành hàng" if n == 0 else
                            f"Chỉ có {n} cột có giá trị (trung vị ngành {tv}) — CMS có thể thiếu thông số"))
    # --- độ phủ cột + cột luôn trống + lệch FILTER/TSKT + gộp nhiều thuộc tính
    spec_dem: Counter = Counter()
    if spec is not None and len(spec):
        skus = set(sku_cate)
        for sku, ma, v in spec[["sku", "ma", "gia_tri"]].itertuples(index=False):
            if sku in skus and ma and v:
                spec_dem[(sku_cate[sku], ma)] += 1
    phu = []
    for cate, b in bang.items():
        n = len(b["rows"])
        nguon = {m: 0 for m in cot_tt(b)}
        for d in (tm.get(cate, {}), fm.get(cate, {})):
            for pid, m in d.items():
                if m in nguon:
                    nguon[m] += 1
        for m in cot_tt(b):
            co = sum(1 for r in b["rows"] if r["vals"].get(m))
            pim = spec_dem.get((cate, m), 0)
            ghi_chu = ""
            if co == 0 and n >= 3:
                ghi_chu = "Không có mapping nào trỏ tới cột này" if not nguon[m] else "Có mapping nhưng CMS không có dữ liệu"
                them(MUC_TB if pim else MUC_THAP, "cot_trong", cate, ma=m,
                     goi_y=ghi_chu + (f" — PIM cũ đang có {pim} SKU có giá trị" if pim else ""))
            phu.append({"Ngành": cate, "Mã cột": m, "Tên cột": b["ten"].get(m, ""), "Loại": "FILTER" if la_cot_filter(m) else "TSKT",
                        "Tool có giá trị": co, "Tổng SKU": n, "% tool": round(100 * co / n, 1) if n else 0,
                        "PIM cũ có": pim, "Số mapping trỏ tới": nguon[m], "Ghi chú": ghi_chu})
        theo_goc = defaultdict(dict)
        for m in cot_tt(b):
            goc = re.sub(r"_(filter|tskt)_master$", "", m)
            theo_goc[goc]["filter" if la_cot_filter(m) else "tskt"] = m
        luon_trong = {m for m in cot_tt(b) if not any(r["vals"].get(m) for r in b["rows"])}
        for goc, cap in theo_goc.items():
            if len(cap) < 2 or luon_trong & set(cap.values()):
                continue
            for r in b["rows"]:
                f, t = bool(r["vals"].get(cap["filter"])), bool(r["vals"].get(cap["tskt"]))
                if f != t:
                    co_cot, thieu = (cap["filter"], cap["tskt"]) if f else (cap["tskt"], cap["filter"])
                    them(MUC_THAP, "lech_ft", cate, ma=thieu, sku=r["sku"],
                         goi_y=f"{co_cot} có giá trị nhưng {thieu} trống")
    for (cate, sku, ma), pids in nguon_cot.items():
        if len(pids) > 1:
            them(MUC_THAP, "gop", cate, pid=", ".join(sorted(pids)), ma=ma, sku=sku,
                 goi_y="Kiểm tra trùng lặp nội dung giữa các thuộc tính CMS gộp chung")
    df = pd.DataFrame(list(nhom.values()))
    if len(df):
        df["SKU ví dụ"] = df["SKU ví dụ"].map(", ".join)
        thu_tu = {MUC_CAO: 0, MUC_TB: 1, MUC_THAP: 2}
        df = df.sort_values(["Mức", "Số SKU"], key=lambda s: s.map(thu_tu) if s.name == "Mức" else -s,
                            ignore_index=True)
    phu_df = pd.DataFrame(phu).sort_values(["Ngành", "% tool"], ignore_index=True) if phu else pd.DataFrame()
    mat_o = int(df.loc[df["_loai"].isin(["filter", "cot", "mat"]) & (df["Mức"] != MUC_THAP), "Số SKU"].sum()) \
        if len(df) else 0
    return {"loi": df, "phu": phu_df, "mat_o": mat_o}


# ============================================================================
# §10 ĐỀ XUẤT SỬA DỮ LIỆU CMS SAI (thành viên đề xuất -> dùng ngay; admin duyệt -> dùng chung lâu dài)
# ============================================================================
# Lưu trữ (không xung đột ghi giữa nhiều người):
#   shared/de_xuat/<tài khoản>.json  — danh sách đề xuất CỦA tài khoản đó (chỉ chủ tài khoản ghi)
#   shared/de_xuat_duyet.json        — quyết định của admin {id: {trang_thai, nguoi_duyet, luc, ghi_chu, gia_tri}}
#   shared/sua_sku.json / sua_gia_tri.json / quy_doi_filter.json — quy tắc ĐÃ DUYỆT, áp cho mọi người
# Đề xuất đang chờ chỉ áp cho workspace của người đề xuất.
DX_CHO, DX_DUYET, DX_TU_CHOI, DX_THU_HOI = "Chờ duyệt", "Đã duyệt", "Từ chối", "Đã thu hồi"
PV_SKU, PV_GIA_TRI, PV_QUY_DOI = "sku", "gia_tri", "quy_doi"
NHAN_PHAM_VI = {PV_SKU: "Chỉ SKU này", PV_GIA_TRI: "Mọi SKU cùng giá trị (trong ngành)",
                PV_QUY_DOI: "Quy đổi FILTER (giá trị CMS → option)"}


def tao_de_xuat(nguoi: str, ten_nguoi: str, pham_vi: str, cate: str, ma: str, gia_tri_cu: str, gia_tri_moi: str,
                sku: str = "", ten_cot: str = "", so_sku: int = 1, ly_do: str = "", sku_vd: str = "") -> dict:
    import secrets
    return {"id": f"{nguoi}-{datetime.now(VN_TZ).strftime('%y%m%d%H%M%S')}-{secrets.token_hex(2)}",
            "nguoi": nguoi, "ten_nguoi": ten_nguoi, "luc": bay_gio(), "pham_vi": pham_vi,
            "loai": "filter" if la_cot_filter(ma) else "tskt", "cate": cate, "sku": sku, "ma": ma,
            "ten_cot": ten_cot, "gia_tri_cu": gia_tri_cu, "gia_tri_moi": gia_tri_moi,
            "so_sku": int(so_sku or 1), "sku_vd": sku_vd or sku, "ly_do": ly_do.strip()}


def khoa_de_xuat(d: dict) -> str:
    """Khoá quy tắc mà đề xuất tạo ra (2 đề xuất cùng khoá = cùng 1 chỗ sửa)."""
    if d["pham_vi"] == PV_SKU:
        return f"{d['sku']}\t{d['ma']}"
    if d["pham_vi"] == PV_QUY_DOI:
        return f"{d['ma']}\t{khoa_quy_doi(d['gia_tri_cu'])}"
    return f"{d['cate']}\t{d['ma']}\t{khoa_quy_doi(d['gia_tri_cu'])}"


def trang_thai_dx(d: dict, duyet: Dict[str, dict]) -> str:
    return (duyet.get(d["id"]) or {}).get("trang_thai", DX_CHO)


def ap_de_xuat_vao_quy_tac(d: dict, sua_sku: dict, sua_gt: dict, quy_doi: dict, gia_tri: Optional[str] = None) -> None:
    """Ghi 1 đề xuất vào bộ quy tắc (dùng cho: duyệt -> bộ chung; đang chờ -> bộ riêng của người đề xuất)."""
    v = d["gia_tri_moi"] if gia_tri is None else gia_tri
    k = khoa_de_xuat(d)
    if d["pham_vi"] == PV_SKU:
        sua_sku[k] = {"cu": khoa_quy_doi(d["gia_tri_cu"]), "moi": v, "id": d["id"]}
    elif d["pham_vi"] == PV_QUY_DOI:
        quy_doi[k] = v
    else:
        sua_gt[k] = v


def bo_quy_tac(d: dict, sua_sku: dict, sua_gt: dict, quy_doi: dict) -> None:
    k = khoa_de_xuat(d)
    {PV_SKU: sua_sku, PV_QUY_DOI: quy_doi}.get(d["pham_vi"], sua_gt).pop(k, None)


def quy_tac_hieu_luc(chung: Tuple[dict, dict, dict], de_xuat_rieng: List[dict],
                     duyet: Dict[str, dict]) -> Tuple[dict, dict, dict, int]:
    """Bộ quy tắc áp cho 1 workspace = quy tắc đã duyệt (chung) + đề xuất CHỜ DUYỆT của chính người đó.
    -> (sua_sku, sua_gt, quy_doi, số đề xuất riêng đang áp)."""
    sk, gt, qd = (dict(x or {}) for x in chung)
    n = 0
    for d in sorted(de_xuat_rieng, key=lambda x: x.get("luc", "")):
        if trang_thai_dx(d, duyet) == DX_CHO:
            ap_de_xuat_vao_quy_tac(d, sk, gt, qd)
            n += 1
    return sk, gt, qd, n


def bang_de_xuat(ds_dx: List[dict], duyet: Dict[str, dict], opt: Optional[dict] = None) -> pd.DataFrame:
    rows = []
    for d in ds_dx:
        q = duyet.get(d["id"]) or {}
        moi = q.get("gia_tri", d["gia_tri_moi"]) if q.get("trang_thai") == DX_DUYET else d["gia_tri_moi"]
        gn = ""
        if opt and d["loai"] == "filter":
            gn = giai_nghia_filter(d["ma"], moi, opt)
        rows.append({"id": d["id"], "Trạng thái": q.get("trang_thai", DX_CHO), "Người đề xuất": d.get("ten_nguoi") or d["nguoi"],
                     "Lúc": d.get("luc", ""), "Phạm vi": NHAN_PHAM_VI.get(d["pham_vi"], d["pham_vi"]),
                     "Ngành": d.get("cate", ""), "SKU": d.get("sku") or d.get("sku_vd", ""), "Mã cột": d["ma"],
                     "Tên cột": d.get("ten_cot", ""), "Giá trị CMS (sai)": d["gia_tri_cu"], "Giá trị đúng": moi,
                     "Giải nghĩa FILTER": gn, "Số SKU": d.get("so_sku", 1), "Lý do": d.get("ly_do", ""),
                     "Admin": q.get("nguoi_duyet", ""), "Lúc duyệt": q.get("luc", ""), "Ghi chú admin": q.get("ghi_chu", "")})
    cols = ["id", "Trạng thái", "Người đề xuất", "Lúc", "Phạm vi", "Ngành", "SKU", "Mã cột", "Tên cột",
            "Giá trị CMS (sai)", "Giá trị đúng", "Giải nghĩa FILTER", "Số SKU", "Lý do", "Admin", "Lúc duyệt",
            "Ghi chú admin"]
    df = pd.DataFrame(rows, columns=cols)
    return df.sort_values("Lúc", ascending=False, ignore_index=True) if len(df) else df


def dem_sku_cung_gia_tri(bang: Dict[str, dict], cate: str, ma: str, gia_tri: str) -> Tuple[int, List[str]]:
    b = bang.get(cate)
    if not b or not gia_tri:
        return 0, []
    k = khoa_quy_doi(gia_tri)
    ds = [r["sku"] for r in b["rows"] if r["vals"].get(ma) and khoa_quy_doi(r["vals"][ma]) == k]
    return len(ds), ds


# ============================================================================
# §11 ĐỘ HOÀN THIỆN (completeness) + QUY TẮC KIỂM TRA (validation rules) — kiểu Akeneo / Salsify / Pimcore
# ============================================================================
# Quy tắc lưu dùng chung: {"<mã ngành hoặc *>\t<mã cột>": {"loai", "tham_so", "muc", "ghi_chu"}}
QT_BAT_BUOC, QT_SO, QT_KHOANG, QT_DANH_SACH, QT_DO_DAI, QT_REGEX = "bat_buoc", "so", "khoang", "danh_sach", "do_dai", "regex"
NHAN_QT = {QT_BAT_BUOC: "Bắt buộc có giá trị", QT_SO: "Là số (+ đơn vị cho phép)", QT_KHOANG: "Số trong khoảng",
           QT_DANH_SACH: "Chỉ nhận giá trị trong danh sách", QT_DO_DAI: "Độ dài tối đa", QT_REGEX: "Khớp mẫu (regex)"}
HUONG_DAN_THAM_SO = {QT_BAT_BUOC: "(để trống)", QT_SO: "đơn vị cho phép, cách nhau dấu phẩy: cm,mm — trống = số trơn",
                     QT_KHOANG: "min-max, vd 0.1-500", QT_DANH_SACH: "các giá trị cách nhau dấu ; vd Có;Không",
                     QT_DO_DAI: "số ký tự tối đa, vd 120", QT_REGEX: r"biểu thức chính quy, vd ^\d+ W$"}
MUC_LOI, MUC_CANH_BAO = "Lỗi", "Cảnh báo"
_SO_VA_DV_RE = re.compile(r"^\s*(-?\d+(?:[.,]\d+)?)\s*([^\d\s].*)?$")


def hang_chu(pct: float) -> str:
    """Hạng A–E như Akeneo Data Quality Insights."""
    return "A" if pct >= 90 else "B" if pct >= 80 else "C" if pct >= 70 else "D" if pct >= 60 else "E"


def quy_tac_cho(quy_tac: Dict[str, dict], cate: str, ma: str) -> List[dict]:
    out = []
    for k in (f"*\t{ma}", f"{cate}\t{ma}"):
        q = quy_tac.get(k)
        if isinstance(q, dict) and q.get("loai"):
            out.append(q)
        elif isinstance(q, list):
            out += [x for x in q if isinstance(x, dict) and x.get("loai")]
    return out


def vi_pham(q: dict, v: str) -> Optional[str]:
    """None = đạt; chuỗi = mô tả vi phạm."""
    loai, ts = q.get("loai"), str(q.get("tham_so") or "").strip()
    if loai == QT_BAT_BUOC:
        return None if v else "Trống — cột bắt buộc"
    if not v:
        return None
    phan = [p.strip() for p in v.split(SEP_TSKT) if p.strip()]
    if loai in (QT_SO, QT_KHOANG):
        dv_cho = [x.strip().lower() for x in ts.split(",") if x.strip()] if loai == QT_SO else []
        for p in phan:
            m = _SO_VA_DV_RE.match(p)
            if not m:
                return f'"{p}" không phải số'
            dv = (m.group(2) or "").strip().lower()
            if loai == QT_SO:
                if dv_cho and dv not in dv_cho:
                    return f'"{p}": đơn vị phải là {", ".join(dv_cho)}'
                if not dv_cho and dv:
                    return f'"{p}": cột chỉ nhận số trơn'
            else:
                try:
                    a, b = [float(x.replace(",", ".")) for x in re.split(r"\s*-\s*(?=\d)", ts, maxsplit=1)]
                except Exception:  # noqa: BLE001
                    return None
                so = float(m.group(1).replace(",", "."))
                if not (a <= so <= b):
                    return f'"{p}" ngoài khoảng {a:g}–{b:g}'
        return None
    if loai == QT_DANH_SACH:
        cho = {khoa_option(x.strip()) for x in ts.split(";") if x.strip()}
        sai = [p for p in phan if khoa_option(p) not in cho]
        return f"Ngoài danh sách: {', '.join(sai[:3])}" if sai and cho else None
    if loai == QT_DO_DAI:
        try:
            n = int(float(ts))
        except Exception:  # noqa: BLE001
            return None
        return f"Dài {len(v)} ký tự > {n}" if len(v) > n else None
    if loai == QT_REGEX and ts:
        try:
            return None if re.fullmatch(ts, v) else f"Không khớp mẫu {ts}"
        except re.error:
            return None
    return None


def do_hoan_thien(bang: Dict[str, dict], cau_hinh: Dict[str, dict], quy_tac: Dict[str, dict],
                  sua: Optional[dict] = None, don_vi: Optional[dict] = None, rong: Optional[dict] = None) -> dict:
    """Độ hoàn thiện từng SKU / ngành (trên GIÁ TRỊ SẼ XUẤT) + vi phạm quy tắc kiểm tra.
    -> {"sku": DataFrame, "nganh": DataFrame, "vi_pham": DataFrame}."""
    sua, don_vi, rong, quy_tac = sua or {}, don_vi or {}, rong or {}, quy_tac or {}
    r_sku, r_ng, r_vp = [], [], []
    for cate, b in bang.items():
        khong_dien = set((cau_hinh.get(cate) or {}).get("cot_khong_dien") or ())
        cot = [m for m in cot_tt(b) if m not in khong_dien]
        qt = {m: quy_tac_cho(quy_tac, cate, m) for m in cot}
        bat_buoc = [m for m in cot if any(q["loai"] == QT_BAT_BUOC for q in qt[m])]
        tong_pct, du_bb = 0.0, 0
        for r in b["rows"]:
            co, thieu = 0, []
            for m in cot:
                v, _ = bien_doi_o(cate, r["sku"], m, r["vals"].get(m, ""), sua, don_vi, rong)
                if v:
                    co += 1
                for q in qt[m]:
                    loi = vi_pham(q, v)
                    if loi:
                        if q["loai"] == QT_BAT_BUOC:
                            thieu.append(m)
                        r_vp.append({"Mức": q.get("muc") or MUC_CANH_BAO, "Ngành": cate, "SKU": r["sku"],
                                     "Mã cột": m, "Tên cột": b["ten"].get(m, ""), "Giá trị": v[:120],
                                     "Quy tắc": NHAN_QT.get(q["loai"], q["loai"]) + (f" ({q.get('tham_so')})"
                                                                                     if q.get("tham_so") else ""),
                                     "Vi phạm": loi, "Ghi chú": q.get("ghi_chu", "")})
            pct = round(100 * co / len(cot), 1) if cot else 0.0
            tong_pct += pct
            du_bb += not thieu
            r_sku.append({"Ngành": cate, "SKU": r["sku"], "Model": r["model"], "% đầy đủ": pct, "Hạng": hang_chu(pct),
                          "Cột có giá trị": co, "Tổng cột": len(cot), "Thiếu cột bắt buộc": ", ".join(thieu),
                          "Đủ bắt buộc": not thieu})
        n = len(b["rows"])
        tb = round(tong_pct / n, 1) if n else 0.0
        r_ng.append({"Ngành": cate, "Tên": b.get("ten_nh", ""), "Số SKU": n, "% đầy đủ TB": tb, "Hạng": hang_chu(tb),
                     "Số cột bắt buộc": len(bat_buoc), "SKU đủ bắt buộc": du_bb,
                     "% SKU sẵn sàng": round(100 * du_bb / n, 1) if n else 0.0})
    cs = ["Ngành", "SKU", "Model", "% đầy đủ", "Hạng", "Cột có giá trị", "Tổng cột", "Thiếu cột bắt buộc", "Đủ bắt buộc"]
    cv = ["Mức", "Ngành", "SKU", "Mã cột", "Tên cột", "Giá trị", "Quy tắc", "Vi phạm", "Ghi chú"]
    return {"sku": pd.DataFrame(r_sku, columns=cs).sort_values(["% đầy đủ", "SKU"], ignore_index=True),
            "nganh": pd.DataFrame(r_ng, columns=["Ngành", "Tên", "Số SKU", "% đầy đủ TB", "Hạng", "Số cột bắt buộc",
                                                 "SKU đủ bắt buộc", "% SKU sẵn sàng"]),
            "vi_pham": pd.DataFrame(r_vp, columns=cv)}


def goi_y_quy_tac(bang: Dict[str, dict], cate: str, toi_thieu: int = 10) -> List[dict]:
    """Tự gợi ý quy tắc từ dữ liệu hiện có (người dùng xem rồi mới lưu):
    cột >= 90% ô là 'số + cùng 1 đơn vị' -> quy tắc SỐ với đơn vị đó; cột có <= 12 giá trị khác nhau -> DANH SÁCH."""
    b = bang.get(cate)
    if not b:
        return []
    out = []
    for m in cot_tt(b):
        if la_cot_filter(m):
            continue
        vs = [r["vals"].get(m, "") for r in b["rows"] if r["vals"].get(m)]
        if len(vs) < toi_thieu:
            continue
        dv = Counter()
        for v in vs:
            for p in v.split(SEP_TSKT):
                mm = _SO_VA_DV_RE.match(p.strip())
                if mm:
                    dv[(mm.group(2) or "").strip().lower()] += 1
        tong = sum(len(v.split(SEP_TSKT)) for v in vs)
        if dv and sum(dv.values()) >= 0.9 * tong:
            chinh = [d for d, n in dv.most_common() if n >= 0.05 * tong]
            out.append({"Mã cột": m, "Tên": b["ten"].get(m, ""), "Loại": QT_SO, "Tham số": ",".join(x for x in chinh if x),
                        "Căn cứ": f"{sum(dv.values())}/{tong} giá trị là số; đơn vị: {dict(dv.most_common(4))}"})
            continue
        khac = Counter(p.strip() for v in vs for p in v.split(SEP_TSKT) if p.strip())
        if 1 < len(khac) <= 12 and len(vs) >= 2 * len(khac):
            out.append({"Mã cột": m, "Tên": b["ten"].get(m, ""), "Loại": QT_DANH_SACH,
                        "Tham số": ";".join(k for k, _ in khac.most_common()),
                        "Căn cứ": f"{len(khac)} giá trị khác nhau trên {len(vs)} SKU"})
    return out


# ============================================================================
# §12 WORKSPACE THEO MẪU (du_lieu_pim.xlsx / "TEST HÀNG LOẠT IMPORT THÔNG SỐ NEW") — đi 2 chiều web <-> desktop
# ============================================================================
EXCEL_MAX_DONG = 1_048_576
TIEU_DE_MAU = {
    "IMPORT": (["model_code", "variant_code", "sku", "category_code"],
               ["Mã model", "Mã biến thể", "Mã sản phẩm ERP", "Mã danh mục PIM"]),
    "DATA PIM": (["Code", "Name", "AttributeTypeName", "AttributeGroupName", "IsActivated", "OptionCode", "OptionValue"],
                 ["Mã thuộc tính", "Tên thuộc tính", "Loại thuộc tính", "Nhóm thuộc tính", "Kích hoạt", "Mã giá trị",
                  "Giá trị"]),
    "MAPPING TSKT MOI": ["MÃ NGÀNH HÀNG CMS", "TÊN NGÀNH HÀNG CMS", "MÃ THUỘC TÍNH TSKT", "TÊN THUỘC TÍNH TSKT",
                         "MÃ MASTER", "TÊN MASTER", "MÃ HỌ", "MÃ THUỘC TÍNH DATA"],
    "MAPPING FILTER MOI": ["MÃ NGÀNH HÀNG CMS", "TÊN NGÀNH HÀNG CMS", "MÃ THUỘC TÍNH FILTER", "TÊN THUỘC TÍNH CMS",
                           "Mã thuộc tính cũ", "Mã thuộc tính mới"],
    "CHỌN NGÀNH HÀNG": ["CHỌN", "MÃ NH", "TÊN NGÀNH HÀNG", "SỐ SKU", "NGUỒN CATE", "TAB TSKT / CẤU HÌNH", "GHI CHÚ"],
}


def gop_chung_theo_nganh(cu: dict, moi: dict) -> Tuple[dict, dict]:
    """Gộp dữ liệu DÙNG CHUNG từ file mẫu vào kho: ngành có trong file -> thay mapping/cấu hình của ngành đó;
    ngành khác giữ nguyên. DATA PIM: mã thuộc tính có trong file -> thay toàn bộ option của mã đó.
    cu/moi: {"cau_hinh", "map_tskt", "map_filter", "data_pim"} -> (kết quả, thống kê)."""
    out, tk = dict(cu), {}
    if moi.get("cau_hinh"):
        ch = dict(cu.get("cau_hinh") or {})
        ch.update(moi["cau_hinh"])
        out["cau_hinh"], tk["cau_hinh"] = ch, f"{len(moi['cau_hinh'])} ngành"
    for k in ("map_tskt", "map_filter"):
        m = moi.get(k)
        if m is not None and len(m):
            c = cu.get(k)
            giu = c[~c.cate.isin(set(m.cate))] if c is not None and len(c) else c
            out[k] = pd.concat([x for x in (giu, m) if x is not None], ignore_index=True)
            tk[k] = f"{m.cate.nunique()} ngành · {len(m):,} dòng"
    m = moi.get("data_pim")
    if m is not None and len(m):
        c = cu.get("data_pim")
        giu = c[~c.Code.isin(set(m.Code))] if c is not None and len(c) else c
        out["data_pim"] = pd.concat([x for x in (giu, m) if x is not None], ignore_index=True)
        tk["data_pim"] = f"{m.Code.nunique():,} mã · {len(m):,} option"
    return out, tk


def xuat_workspace_mau(imp: pd.DataFrame, data_sp: pd.DataFrame, data_pim: pd.DataFrame, cau_hinh: Dict[str, dict],
                       map_tskt: pd.DataFrame, map_filter: pd.DataFrame, bang: Dict[str, dict],
                       chon: Optional[List[dict]] = None, log: Optional[List[list]] = None,
                       sua: Optional[dict] = None, don_vi: Optional[dict] = None, rong: Optional[dict] = None,
                       chi_nganh_trong_lo: bool = True) -> bytes:
    """Ghi toàn bộ workspace ra 1 file Excel ĐÚNG BỐ CỤC MẪU (mở bằng Excel / bản desktop 66.py được).
    Tab "TSKT <mã> <tên>" chứa GIÁ TRỊ SẼ XUẤT (đã áp sửa tay / đơn vị / quy tắc Không-Đang cập nhật)."""
    import xlsxwriter
    sua, don_vi, rong = sua or {}, don_vi or {}, rong or {}
    nganh = set(bang) if (chi_nganh_trong_lo and bang) else None
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True, "strings_to_numbers": False, "strings_to_formulas": False,
                                   "strings_to_urls": False})
    f_txt = wb.add_format({"num_format": "@"})
    f_dam = wb.add_format({"num_format": "@", "bold": True})

    def sheet(ten, header_rows, rows, rong_cot=None):
        # Excel tối đa 1.048.576 dòng/sheet: dữ liệu dài hơn tự tách sang "<tên> (2)", "<tên> (3)"… (đọc lại tự gộp)
        goc_ten = ten_tab_desktop(ten) if ten.upper().startswith("TSKT ") else ten
        n = max([len(r) for r in header_rows] + [7])
        so_ws, i, ws = 0, EXCEL_MAX_DONG, None
        for r in rows:
            if i >= EXCEL_MAX_DONG:
                so_ws += 1
                ws = wb.add_worksheet(goc_ten if so_ws == 1 else f"{goc_ten[:25]} ({so_ws})")
                ws.set_column(0, n - 1, rong_cot or 18, f_txt)
                i = 0
                for h in header_rows:
                    for j, v in enumerate(h):
                        if v not in (None, ""):
                            ws.write_string(i, j, str(v), f_dam)
                    i += 1
                ws.freeze_panes(len(header_rows), 0)
            for j, v in enumerate(r):
                if isinstance(v, bool):
                    ws.write_boolean(i, j, v)
                elif v not in (None, ""):
                    ws.write_string(i, j, str(v), f_txt)
            i += 1
        if ws is None:  # không có dòng dữ liệu -> vẫn tạo sheet có tiêu đề
            ws = wb.add_worksheet(goc_ten)
            ws.set_column(0, n - 1, rong_cot or 18, f_txt)
            for k, h in enumerate(header_rows):
                for j, v in enumerate(h):
                    if v not in (None, ""):
                        ws.write_string(k, j, str(v), f_dam)
            ws.freeze_panes(len(header_rows), 0)
        return ws
    h1, h2 = TIEU_DE_MAU["IMPORT"]
    sheet("IMPORT", [h1, h2], ([r.model_code, r.variant_code, r.sku, r.category_code]
                               for r in imp.itertuples(index=False)))
    sheet("DATA SP", [COT_DATA_SP], data_sp[COT_DATA_SP].itertuples(index=False, name=None) if len(data_sp) else [])
    h1, h2 = TIEU_DE_MAU["DATA PIM"]
    sheet("DATA PIM", [h1, h2], data_pim[COT_DATA_PIM].itertuples(index=False, name=None) if len(data_pim) else [])
    ch_rows = []
    for cate, cfg in cau_hinh.items():
        if nganh is not None and cate not in nganh:
            continue
        kd = set(cfg.get("cot_khong_dien") or ())
        cot = [c for c in cfg.get("cot", []) if c not in kd]
        ch_rows.append([cate, cfg.get("ten", "")] + cot)
        ch_rows.append(["", ""] + [cfg.get("ten_cot", {}).get(c, "") for c in cot])
    sheet("CẤU HÌNH CATEGORY", [["CATEGORY ID", "CATEGORY NAME", "DANH SÁCH TSKT (ngang, mỗi ô 1 mã)"]], ch_rows)
    mt = map_tskt if nganh is None else map_tskt[map_tskt.cate.isin(nganh)]
    ho = {c: (cfg.get("ma_ho") or "") for c, cfg in cau_hinh.items()}
    sheet("MAPPING TSKT MOI", [TIEU_DE_MAU["MAPPING TSKT MOI"]],
          ([r.cate, r.cate_name, r.prop_id, r.prop_name, r.ma, r.ten_ma, ho.get(r.cate, ""), ""]
           for r in mt.itertuples(index=False)))
    mf = map_filter if nganh is None else map_filter[map_filter.cate.isin(nganh)]
    sheet("MAPPING FILTER MOI", [TIEU_DE_MAU["MAPPING FILTER MOI"]],
          ([r.cate, r.cate_name, r.prop_id, r.prop_name, r.ma, r.ma] for r in mf.itertuples(index=False)))
    hc = TIEU_DE_MAU["CHỌN NGÀNH HÀNG"]
    sheet("CHỌN NGÀNH HÀNG", [hc], ([x.get(k, "") for k in hc] for x in (chon or [])))
    for cate, b in bang.items():
        cot = b["attr"]
        rows = []
        for r in b["rows"]:
            out = [r["vals"].get("model_code", r["model"]), r["sku"], r["vals"].get("variant_code", r["variant"])]
            for c in cot:
                out.append(bien_doi_o(cate, r["sku"], c, r["vals"].get(c, ""), sua, don_vi, rong)[0])
            rows.append(out)
        sheet(b["title"], [["model_code", "sku", "variant_code"] + cot,
                           ["Mã model", "Mã sản phẩm ERP", "Mã biến thể"] + [b["ten"].get(c, "") for c in cot]], rows)
    if log:
        sheet("LOG", [["THỜI GIAN", "SKU", "CATEGORY", "TSKT", "NGUYÊN NHÂN"]],
              ([bay_gio()] + list(x)[:4] for x in log))
    wb.close()
    return buf.getvalue()



# ============================================================================
# §13 NẠP NHANH 1 CỤC (tự nhận loại file) + TỰ LỌC + FILE XIN DATA CMS
# ============================================================================
def doc_mau_nganh(rows: List[list]) -> Dict[str, dict]:
    """File "Export Product Template" của 1 ngành (PIM): dòng 1 = [tên ngành | mã ngành | mã cột thuộc tính…],
    dòng 2 = [trống | trống | tên tiếng Việt từng cột]. Không có dòng dữ liệu. -> cấu hình ngành
    {mã ngành: {"ten", "cot", "ten_cot"}} (giống CẤU HÌNH CATEGORY). Không đúng bố cục -> {}."""
    if not rows or len(rows) > 3:
        return {}
    r1 = [chuan_hoa_key(x) for x in rows[0]]
    if len(r1) < 4:
        return {}
    a, b = r1[0], r1[1]
    id_a, id_b = chuan_hoa_id(a), chuan_hoa_id(b)
    if re.fullmatch(r"\d+", id_b or "") and a and not re.fullmatch(r"\d+", id_a or ""):
        cate, ten = id_b, a
    elif re.fullmatch(r"\d+", id_a or "") and b and not re.fullmatch(r"\d+", id_b or ""):
        cate, ten = id_a, b
    else:
        return {}
    ma = [(j, v) for j, v in enumerate(r1) if j >= 2 and v]
    if len(ma) < 2 or not all(re.fullmatch(r"[A-Za-z0-9_]+", v) for _, v in ma) or not any("_" in v for _, v in ma):
        return {}
    r2 = [chuan_hoa_key(x) for x in rows[1]] if len(rows) > 1 else []
    if len(rows) > 2 and any(chuan_hoa_key(x) for x in rows[2]):
        return {}  # có dòng dữ liệu -> không phải file mẫu trống
    cot, ten_cot = [], {}
    for j, v in ma:
        # LOẠI 8 cột cố định (= COT_KHONG_PHAI_SPEC): model_code, sku, category_code, variant_code,
        # family_code, family_variant_code, model_activated, variant_activated
        if v.lower() in COT_KHONG_PHAI_SPEC or v in cot:
            continue
        cot.append(v)
        if j < len(r2) and r2[j]:
            ten_cot[v] = r2[j]
    return {cate: {"ten": ten, "cot": cot, "ten_cot": ten_cot}} if cot else {}




def doc_cot_tu_sku(rows: List[list]) -> Tuple[str, List[Tuple[str, str]]]:
    """Trích danh sách cột THUỘC TÍNH từ file export PIM (có dữ liệu SKU): dòng 1 = mã cột, dòng 2 = tên.
    Trả về (cate_id gợi ý từ category_code của dữ liệu, [(ma_cot, ten_vn), ...]).
    Loại 8 cột cố định (COT_KHONG_PHAI_SPEC). Dùng để OFFER thêm cột vào Cấu hình ngành."""
    if not rows or len(rows) < 2:
        return "", []
    h1 = [chuan_hoa_key(h) for h in rows[0]]
    h2 = [chuan_hoa_key(h) for h in rows[1]] if len(rows) > 1 else []
    cot = []
    for j, v in enumerate(h1):
        if not v or v.lower() in COT_KHONG_PHAI_SPEC:
            continue
        if not re.fullmatch(r"[A-Za-z0-9_]+", v) or "_" not in v:
            continue
        ten = h2[j] if j < len(h2) else ""
        cot.append((v, ten))
    # tìm category_code có mặt
    idx = {h.lower(): j for j, h in enumerate(h1)}
    j_cate = idx.get("category_code")
    if j_cate is None:
        j_cate = tim_cot_truong(h1, "category_code")
    cate_id = ""
    if j_cate is not None and j_cate >= 0:
        from collections import Counter as _C
        c = _C()
        for r in rows[2:]:
            if j_cate < len(r):
                v = chuan_hoa_id(r[j_cate])
                if v:
                    c[v] += 1
        if c:
            cate_id = c.most_common(1)[0][0]
    return cate_id, cot


LOAI_FILE = {"nganh": "File mẫu ngành hàng (Export Product Template) → Cấu hình ngành",
             "mau": "Workspace theo mẫu (nhiều sheet)", "cms": "File CMS export (DATA SP)",
             "sku": "Danh sách SKU / file export PIM (model, SKU, biến thể ± TSKT)"}


def nhan_dien_file(data: bytes, ten_file: str) -> dict:
    """Đoán loại file + đọc luôn. -> {"loai", "mo_ta", ...dữ liệu đã đọc}."""
    so = SoExcel(data, ten_file)
    ten = so.sheets
    if {"IMPORT", "DATA SP"} & set(ten) and len(ten) > 1 or "MAPPING TSKT MOI" in ten:
        w = doc_workspace_cu(data, ten_file)
        return {"loai": "mau", **w}
    cfg_ng: Dict[str, dict] = {}
    for sh in ten:
        cfg_ng.update(doc_mau_nganh(so.rows(sh)))
    if cfg_ng:
        return {"loai": "nganh", "cau_hinh": cfg_ng}
    for sh in ten:
        rows = so.rows(sh)
        if not rows:
            continue
        h = [ep_text(x) for x in rows[0]]
        if tim_cot_truong(h, "PROPERTYID") >= 0 and tim_cot_truong(h, "PROPVALUE") >= 0:
            df, loi = doc_cms_export(data, ten_file)
            return {"loai": "cms", "data_sp": df, "loi": loi, "sheet": sh}
        n = max(len(x) for x in rows[:50])
        if 6 <= n <= 8 and len(rows) > 2:  # CMS export KHÔNG tiêu đề / tiêu đề lạ: cột D số (PROPERTYID), F giá trị
            mau = rows[1:200]
            so_d = sum(bool(re.fullmatch(r"\d+", chuan_hoa_id(x[3]))) for x in mau if len(x) > 5)
            co_f = sum(bool(chuan_hoa_key(x[5])) for x in mau if len(x) > 5)
            if so_d >= 0.8 * len(mau) and co_f >= 0.5 * len(mau):
                df, loi = doc_cms_export(data, ten_file)
                if not loi:
                    return {"loai": "cms", "data_sp": df, "loi": None, "sheet": sh,
                            "ghi_chu": ["Không nhận ra tiêu đề — đọc theo vị trí cột như desktop (A..G)."]}
        r = doc_mot_cuc(rows)
        if len(r["import"]):
            # Trích danh sách cột thuộc tính (gợi ý cấu hình ngành) từ chính file này.
            cate_gy, cot_gy = doc_cot_tu_sku(rows)
            r["goi_y_cfg"] = {"cate_id": cate_gy, "cot": cot_gy}
            return {"loai": "sku", **r, "sheet": sh}
    return {"loai": None, "loi": "Không nhận ra loại file (cần: workspace mẫu / CMS export / danh sách SKU)."}


def loc_data_sp(data_sp: pd.DataFrame, imp: pd.DataFrame) -> Tuple[pd.DataFrame, dict]:
    """Chỉ giữ dòng DATA SP của SKU có trong IMPORT (như bước 'nạp và lọc' của desktop) — nhẹ, nhanh hơn."""
    if data_sp is None or not len(data_sp) or imp is None or not len(imp):
        return data_sp, {"giu": len(data_sp) if data_sp is not None else 0, "bo": 0, "sku_bo": 0}
    keep = data_sp.PRODUCTCODE.isin(set(imp.sku))
    tk = {"giu": int(keep.sum()), "bo": int((~keep).sum()),
          "sku_bo": int(data_sp.loc[~keep, "PRODUCTCODE"].nunique())}
    return data_sp[keep].reset_index(drop=True), tk


def ds_xin_data(imp: pd.DataFrame, data_sp: pd.DataFrame, bang: Dict[str, dict],
                log: Optional[List[list]] = None) -> Dict[str, pd.DataFrame]:
    """ERP (SKU) / MODEL KHÔNG CÓ GIÁ TRỊ -> danh sách đi xin data CMS.
    Lý do: CMS chưa có dữ liệu · ngành chưa có mapping/cấu hình · có dữ liệu nhưng không map được thuộc tính nào."""
    cols = ["model_code", "sku", "variant_code", "category_code", "CATEGORYID (CMS)", "Tên sản phẩm", "Lý do",
            "Số dòng CMS", "Số ô map được"]
    if imp is None or not len(imp):
        return {"sku": pd.DataFrame(columns=cols), "model": pd.DataFrame()}
    sp = data_sp if data_sp is not None else pd.DataFrame(columns=COT_DATA_SP)
    dem = Counter(sp.PRODUCTCODE) if len(sp) else Counter()
    ten_sp, cate_sp = {}, {}
    if len(sp):
        for a, b, c in sp[["PRODUCTCODE", "PRODUCTNAME", "CATEGORYID"]].drop_duplicates("PRODUCTCODE").itertuples(
                index=False):
            ten_sp[a], cate_sp[a] = b, c
    o_map = {}
    for b in bang.values():
        ctt = cot_tt(b)
        for r in b["rows"]:
            o_map[r["sku"]] = sum(1 for m in ctt if r["vals"].get(m))
    ly_do_log = {}
    for x in log or []:
        if x and x[0] and len(x) > 3 and x[0] not in ly_do_log and (
                "mapping" in str(x[3]) or "cấu hình" in str(x[3]) or "DATA SP" in str(x[3])):
            ly_do_log[x[0]] = str(x[3])
    cate_bo = {x[1] for x in log or [] if x and not x[0] and len(x) > 3 and "cấu hình" in str(x[3])}
    out = []
    for r in imp.itertuples(index=False):
        sku = r.sku
        if not dem.get(sku):
            ly = "CMS chưa có dữ liệu (SKU không có trong DATA SP)"
        elif sku not in o_map:
            ly = ly_do_log.get(sku) or ("Ngành chưa có tab TSKT / cấu hình" if cate_sp.get(sku) in cate_bo
                                        else "Ngành chưa có mapping")
        elif o_map[sku] == 0:
            ly = "Có dữ liệu CMS nhưng KHÔNG thuộc tính nào map được (kiểm tra mã thuộc tính/ngành)"
        else:
            continue
        out.append([r.model_code, sku, r.variant_code, r.category_code, cate_sp.get(sku, ""), ten_sp.get(sku, ""), ly,
                    dem.get(sku, 0), o_map.get(sku, 0)])
    df = pd.DataFrame(out, columns=cols)
    # theo MODEL: model mà MỌI SKU đều không có giá trị
    tong = Counter(imp.model_code[imp.model_code != ""])
    thieu = Counter(df.model_code[df.model_code != ""]) if len(df) else Counter()
    m = [{"model_code": k, "Số SKU": tong[k], "Số SKU không có giá trị": v,
          "Tình trạng": "MODEL KHÔNG CÓ GIÁ TRỊ (mọi SKU)" if v == tong[k] else "Một phần SKU thiếu"}
         for k, v in thieu.items()]
    dm = pd.DataFrame(m, columns=["model_code", "Số SKU", "Số SKU không có giá trị", "Tình trạng"])
    if len(dm):
        dm = dm.sort_values(["Tình trạng", "Số SKU"], ascending=[False, False], ignore_index=True)
    return {"sku": df, "model": dm}
